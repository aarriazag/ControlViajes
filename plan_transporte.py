# -*- coding: utf-8 -*-
"""
PLAN DE TRANSPORTE — carga de planes por Excel, asignación de camiones y envío a SimpliRoute (SR).

Qué hace (todo vive aquí; app.py solo lo enchufa con unas pocas líneas):
  1. Lee la plantilla de carga masiva del portal de pedidos (hoja PEDIDOS) y la VALIDA fila por fila.
  2. Agrupa: cada fila es una PARADA; las filas con el mismo Cliente Orden + Fecha + G. Ruta son UN viaje
     (una ruta en SR con una visita por parada). G. Ruta en blanco = la fila es su propio viaje.
  3. Guarda todo como borrador en Supabase ANTES de tocar SR (si algo falla, nada se pierde).
  4. Crea en SR el plan GT_<CLIENTE>_<FECHA>, una ruta por viaje (con el vehículo dummy del cliente) y una
     visita por parada, por lotes y sin duplicar nada si se repite el envío.
  5. Asigna placa y piloto desde Control de Ruta (PATCH a la ruta de SR); lo que se cambie en SR se detecta.
  6. Seguimiento: solicitados vs planeados vs con camión.

Principios (iguales al resto de la integración con SR):
  - Nada de aquí tumba la app: las funciones devuelven (ok, mensaje) o un dict de resultado.
  - Nada de SR se guarda en memoria de sesión: solo ids y estados, en Supabase.
  - SR no permite borrar un plan por API: el plan se crea una sola vez por nombre y se reutiliza.
"""
import io
import re
import time as _time
import unicodedata
import warnings
from contextlib import closing
from datetime import date, datetime, time as dtime, timedelta

import integracion_simpliroute as sr_int

# ---------------------------------------------------------------------------
# Interruptores y límites (se cambian aquí, sin tocar el resto)
# ---------------------------------------------------------------------------
PLAN_TRANSPORTE_ACTIVO = True      # False = la página avisa que el módulo está apagado
ENVIO_A_SR_ACTIVO = True           # False = todo funciona como borrador, sin enviar nada a SR
MODO_PRUEBA_POR_DEFECTO = True     # True = los planes nuevos llevan " -PRUEBA" en el nombre hasta que se desmarque
LOTE_ENVIO = 25                    # máximo de rutas+visitas que se envían a SR por clic
DIAS_RETENCION = 90                # planes cerrados con más días que esto se pueden purgar
MAX_FILAS_EXCEL = 500
DIAS_COORD_SR = 3                  # días de visitas de SR que se revisan para completar coordenadas que falten
MAX_VIAJES_POR_PLACA_DIA = 2       # un camión REAL en más viajes el mismo día casi seguro es un error de la plantilla (doble viaje = 2)
PERFILES_PLAN = ("Administrador", "SuperAdministrador", "Operador", "Supervisor")
PERFILES_CONFIG = ("Administrador", "SuperAdministrador")
HORA_INICIO_RUTA = "07:00:00"
HORA_FIN_RUTA = "17:00:00"
MARCA_PRUEBA = "PRUEBA"
DUMMY_GENERICO = "ADICIONALES"     # vehículos genéricos de SR (ADICIONALES1 SECO, ADICIONALES2 SECO…): los usa todo cliente sin dummies propios

# Listas de la hoja PARAMETROS de la plantilla (solo como ayuda en pantalla; la validación real
# usa lo que esté configurado en Control de Ruta).
CLIENTES_ORDEN_PARAMETROS = [
    "GT-CUENTAS VARIAS", "GT-ROPECOR", "GT-NATURACEITES LOCAL", "GT-NATURACEITES QUETZALTENANGO",
    "GT-NATURACEITES TECULUM", "GT-DOS PINOS", "GT-DINANT", "GT-DOLLAR CITY", "GT-HERBALIFE",
    "GT-NESTLE GUATEMALA QUETZALTENANGO", "GT-NESTLE GUATEMALA LOCAL", "GT-NESTLE GUATEMALA TECULUTAN",
    "GT-KELLOGG", "GT-INASA", "GT-BAYER", "GT-CHOCOLATE RICO", "GT-PETOY", "GT-CALSA", "GT-BOQUITAS, S.A",
    "GT-PATSY,S.A", "GT-ESPECIAS DON JULIO", "GT-PRICESMART S.A.", "GT-OPG", "GT-UNIVAR SOLUTION GUATEMALA",
    "GT-DINAMO", "GT-TRASLADOS INTERNOS RANSA", "GT-HEALTHY GREEN WORLD GUATEMALA ,SOCIEDAD ANÓNIMA",
    "GT-GRUPO BRIO", "GT-SIGMAQ Q.(Litozadik, s.a)", "GT-INCAP", "GT-CORDIALSA", "GT-YA ESTA", "GT-TOTTO",
    "GT-LA COMERCIAL SOCIEDAD ANONIMA", "GT-DIANA", "GT-AJEMAYA",
]
DEPOSITOS_PARAMETROS = [
    "GT - DOLLAR CITY MICHATOYA", "GT - CD RANSA QUETZALTENANGO", "GT - CD RANSA TECULUTAN",
    "GT - CD ARCHIVO NARANJO 2", "GT-CD_VILLA_NUEVA_SECO", "GT - CD ARCHIVO NARANJO 1",
    "GT-CD_AMATITLAN_FRIO", "GT-CD_AMATITLAN_SECO", "GT - CD ARCHIVO SAN CRISTOBAL",
]

# Destinos de prueba (dummies). Códigos altos (9xxx) para que no choquen con códigos reales de tienda.
DESTINOS_PRUEBA = {
    "GT-DOLLAR CITY": [
        ("Dollarcity 1", 9001, 14.634900, -90.506900, "Zona 1, Ciudad de Guatemala"),
        ("Dollarcity 2", 9002, 14.599500, -90.513300, "Zona 9, Ciudad de Guatemala"),
        ("Dollarcity 3", 9003, 14.640700, -90.553000, "Zona 4, Mixco"),
        ("Dollarcity 4", 9004, 14.580000, -90.550000, "Zona 11, Ciudad de Guatemala"),
        ("Dollarcity 5", 9005, 14.610000, -90.520000, "Zona 10, Ciudad de Guatemala"),
    ],
    "GT-DIANA": [
        ("Diana 1", 9101, 14.630000, -90.550000, "Zona 7, Ciudad de Guatemala"),
        ("Diana 2", 9102, 14.570000, -90.530000, "Zona 12, Ciudad de Guatemala"),
        ("Diana 3", 9103, 14.660000, -90.510000, "Zona 18, Ciudad de Guatemala"),
    ],
}

_ESQUEMA_LISTO = False


# ---------------------------------------------------------------------------
# Utilidades pequeñas
# ---------------------------------------------------------------------------
def hoy_guatemala():
    return sr_int._hoy_guatemala()


def _norm(s):
    """Minúsculas, sin tildes y con espacios simples — para comparar nombres."""
    s = unicodedata.normalize("NFD", str(s if s is not None else ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip().lower()


def _clave_cliente(s):
    """Clave para comparar 'GT-DIANA', 'gt diana' y 'DIANA' como lo mismo."""
    n = _norm(s)
    n = re.sub(r"^gt\s*[-_ ]\s*", "", n)
    return n


def _txt(v):
    if v is None or (isinstance(v, float) and v != v):      # None o NaN (así llegan las celdas vacías al leer un Excel con pandas)
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _prefijo_codigo(cliente_orden):
    """'GT-DIANA' -> 'DIANA' (el No. de pedido es Cliente + fecha + consecutivo)."""
    return re.sub(r"^GT\s*[-_ ]\s*", "", str(cliente_orden).strip(), flags=re.IGNORECASE).strip().upper()


def _slug_plan(cliente_orden):
    base = _prefijo_codigo(cliente_orden)
    return re.sub(r"[^A-Z0-9]+", "_", base).strip("_") or "CLIENTE"


def nombre_plan(cliente_orden, fecha, prueba=False):
    """GT_<CLIENTE>_<AAAA-MM-DD>, y ' -PRUEBA' al final en modo prueba."""
    n = f"GT_{_slug_plan(cliente_orden)}_{fecha:%Y-%m-%d}"
    return n + (f" -{MARCA_PRUEBA}" if prueba else "")


def codigo_pedido(cliente_orden, fecha, consecutivo):
    return f"{_prefijo_codigo(cliente_orden)} {fecha:%d%m%Y}-{consecutivo:02d}"


def _corto(s, n=300):
    s = "" if s is None else str(s)
    return s if len(s) <= n else s[:n] + "…"


# ---------------------------------------------------------------------------
# Esquema (todo IF NOT EXISTS: en una base que ya lo tiene no cambia nada)
# ---------------------------------------------------------------------------
def asegurar_esquema_plan(conn):
    global _ESQUEMA_LISTO
    with conn.cursor() as cur:
        # Columnas que ya añade la integración con SR; se repiten por si el orden de arranque cambia.
        cur.execute("ALTER TABLE cat_clientes_tiendas ADD COLUMN IF NOT EXISTS latitud DOUBLE PRECISION")
        cur.execute("ALTER TABLE cat_clientes_tiendas ADD COLUMN IF NOT EXISTS longitud DOUBLE PRECISION")
        cur.execute("ALTER TABLE cat_clientes_tiendas ADD COLUMN IF NOT EXISTS direccion TEXT")
        # Un código identifica UNA tienda dentro de cada cliente (el mismo código en otro cliente no choca). Si ya hubiera repetidos,
        # el índice no se crea (la app sigue igual) y se pueden ver en Configuración → Mantenimiento.
        try:
            cur.execute("SAVEPOINT sp_idx_codigo")
            cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS uq_tiendas_cliente_codigo ON cat_clientes_tiendas (cliente, codigo_tienda) "
                        "WHERE codigo_tienda IS NOT NULL")
            cur.execute("RELEASE SAVEPOINT sp_idx_codigo")
        except Exception:
            cur.execute("ROLLBACK TO SAVEPOINT sp_idx_codigo")
        cur.execute("CREATE TABLE IF NOT EXISTS cat_mapeo_cliente_sr (visit_type TEXT PRIMARY KEY, cliente TEXT)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS cat_transporte_cliente (
                cliente_orden TEXT PRIMARY KEY,
                cliente TEXT NOT NULL,
                visit_type TEXT,
                vehiculo_dummy_sr BIGINT,
                vehiculo_dummy_nombre TEXT,
                deposito_nombre TEXT,
                deposito_lat DOUBLE PRECISION,
                deposito_lon DOUBLE PRECISION,
                activo BOOLEAN NOT NULL DEFAULT TRUE
            )""")
        cur.execute("ALTER TABLE cat_transporte_cliente ADD COLUMN IF NOT EXISTS dummy_prefijo TEXT")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS plan_transporte (
                id SERIAL PRIMARY KEY,
                cliente_orden TEXT NOT NULL,
                cliente TEXT NOT NULL,
                fecha DATE NOT NULL,
                nombre TEXT NOT NULL UNIQUE,
                plan_id_sr TEXT,
                prueba BOOLEAN NOT NULL DEFAULT FALSE,
                creado_por TEXT,
                creado TIMESTAMP DEFAULT NOW()
            )""")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS plan_viaje (
                id SERIAL PRIMARY KEY,
                plan_id INTEGER NOT NULL REFERENCES plan_transporte(id) ON DELETE CASCADE,
                num INTEGER NOT NULL,
                grupo INTEGER,
                estado TEXT NOT NULL DEFAULT 'pendiente',
                route_id_sr TEXT,
                placa TEXT, piloto TEXT, vehiculo_sr_id BIGINT, piloto_sr_id BIGINT,
                deposito_nombre TEXT, deposito_lat DOUBLE PRECISION, deposito_lon DOUBLE PRECISION,
                id_viaje_cr TEXT,
                cambiado_en_sr BOOLEAN NOT NULL DEFAULT FALSE,
                error TEXT,
                asignado_por TEXT, asignado_en TIMESTAMP,
                creado TIMESTAMP DEFAULT NOW(),
                UNIQUE (plan_id, num)
            )""")
        cur.execute("ALTER TABLE plan_viaje ADD COLUMN IF NOT EXISTS dummy_sr_id BIGINT")
        # Un solo viaje ABIERTO por grupo dentro de un plan (el candado contra duplicar un viaje por G. Ruta).
        cur.execute("""
            CREATE UNIQUE INDEX IF NOT EXISTS uq_plan_viaje_grupo_abierto
            ON plan_viaje (plan_id, grupo)
            WHERE grupo IS NOT NULL AND estado NOT IN ('cancelado', 'despachado')""")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS plan_parada (
                id SERIAL PRIMARY KEY,
                viaje_id INTEGER NOT NULL REFERENCES plan_viaje(id) ON DELETE CASCADE,
                plan_id INTEGER NOT NULL REFERENCES plan_transporte(id) ON DELETE CASCADE,
                orden INTEGER NOT NULL,
                consecutivo INTEGER NOT NULL,
                codigo_pedido TEXT NOT NULL,
                cod_destino TEXT,
                destino TEXT NOT NULL,
                direccion TEXT,
                latitud DOUBLE PRECISION, longitud DOUBLE PRECISION,
                ventana_ini TEXT, ventana_fin TEXT, duracion TEXT,
                observaciones TEXT, ref_adicional TEXT,
                estado TEXT NOT NULL DEFAULT 'pendiente',
                visit_id_sr BIGINT,
                error TEXT,
                creado TIMESTAMP DEFAULT NOW(),
                UNIQUE (plan_id, codigo_pedido)
            )""")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS plan_carga (
                id SERIAL PRIMARY KEY,
                momento TIMESTAMP DEFAULT NOW(),
                usuario TEXT, archivo TEXT, prueba BOOLEAN NOT NULL DEFAULT FALSE, filas INTEGER
            )""")
        cur.execute("ALTER TABLE plan_parada ADD COLUMN IF NOT EXISTS carga_id INTEGER")
        cur.execute("CREATE INDEX IF NOT EXISTS ix_plan_parada_carga ON plan_parada (carga_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS ix_plan_viaje_plan ON plan_viaje (plan_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS ix_plan_parada_viaje ON plan_parada (viaje_id)")
        cur.execute("CREATE INDEX IF NOT EXISTS ix_plan_transporte_fecha ON plan_transporte (fecha, cliente_orden)")
        cur.execute("""
            CREATE TABLE IF NOT EXISTS plan_bitacora (
                id SERIAL PRIMARY KEY,
                momento TIMESTAMP DEFAULT NOW(),
                usuario TEXT, accion TEXT, plan_id INTEGER, viaje_id INTEGER, detalle TEXT
            )""")
        cur.execute("CREATE INDEX IF NOT EXISTS ix_plan_bitacora_plan ON plan_bitacora (plan_id, momento)")
    conn.commit()
    _ESQUEMA_LISTO = True


def asegurar_esquema_una_vez(get_conn_fn):
    """Para la página: crea el esquema la primera vez que se entra en este proceso. Nunca lanza."""
    if _ESQUEMA_LISTO:
        return True
    try:
        with closing(get_conn_fn()) as conn:
            asegurar_esquema_plan(conn)
        return True
    except Exception:
        return False


def _bitacora(cur, usuario, accion, plan_id=None, viaje_id=None, detalle=""):
    cur.execute(
        "INSERT INTO plan_bitacora (usuario, accion, plan_id, viaje_id, detalle) VALUES (%s,%s,%s,%s,%s)",
        (usuario, accion, plan_id, viaje_id, _corto(detalle, 800)))


# ---------------------------------------------------------------------------
# Configuración por cliente (Cliente Orden -> cliente de Control de Ruta, dummy, depósito, visit_type)
# ---------------------------------------------------------------------------
def listar_config(get_conn_fn, solo_activos=False):
    import psycopg2.extras
    with closing(get_conn_fn()) as conn, conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute("SELECT * FROM cat_transporte_cliente" + (" WHERE activo" if solo_activos else "") +
                    " ORDER BY cliente_orden")
        return [dict(r) for r in cur.fetchall()]


def clientes_transporte(get_conn_fn, permitidos=None):
    """Clientes de Transporte configurados y activos que este usuario puede ver. permitidos=None => todos."""
    cfgs = listar_config(get_conn_fn, solo_activos=True)
    if permitidos is None:
        return cfgs
    ok = set(permitidos)
    return [c for c in cfgs if c["cliente"] in ok]


def hay_clientes_transporte(get_conn_fn):
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT 1 FROM cat_transporte_cliente WHERE activo LIMIT 1")
        return cur.fetchone() is not None


def tiene_acceso(get_conn_fn, perfil, clientes_permitidos):
    """¿Este usuario puede entrar al Plan de Transporte? Solo si su perfil lo permite Y tiene acceso a al menos un cliente de Transporte
    configurado. Única excepción: un Administrador, mientras NO exista ningún cliente de Transporte configurado en todo el sistema
    (para poder hacer la primera configuración). Nunca lanza."""
    if not PLAN_TRANSPORTE_ACTIVO or perfil not in PERFILES_PLAN:
        return False
    try:
        if clientes_transporte(get_conn_fn, clientes_permitidos):
            return True
        return perfil in PERFILES_CONFIG and not hay_clientes_transporte(get_conn_fn)
    except Exception:
        return False


def guardar_config(get_conn_fn, datos, usuario):
    """Crea o actualiza la configuración de un Cliente Orden. Devuelve (ok, mensaje)."""
    co = _txt(datos.get("cliente_orden"))
    cli = _txt(datos.get("cliente"))
    if not co or not cli:
        return False, "Falta el Cliente Orden o el cliente de Control de Ruta."
    vt = _txt(datos.get("visit_type")) or None
    try:
        with closing(get_conn_fn()) as conn:
            try:
                with conn.cursor() as cur:
                    cur.execute("SELECT 1 FROM cat_clientes WHERE nombre = %s", (cli,))
                    if not cur.fetchone():
                        return False, f"El cliente '{cli}' no existe en Control de Ruta (Catálogos → Clientes)."
                    cur.execute("""
                        INSERT INTO cat_transporte_cliente
                          (cliente_orden, cliente, visit_type, vehiculo_dummy_sr, vehiculo_dummy_nombre,
                           deposito_nombre, deposito_lat, deposito_lon, activo, dummy_prefijo)
                        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                        ON CONFLICT (cliente_orden) DO UPDATE SET
                          cliente=EXCLUDED.cliente, visit_type=EXCLUDED.visit_type,
                          vehiculo_dummy_sr=EXCLUDED.vehiculo_dummy_sr,
                          vehiculo_dummy_nombre=EXCLUDED.vehiculo_dummy_nombre,
                          deposito_nombre=EXCLUDED.deposito_nombre, deposito_lat=EXCLUDED.deposito_lat,
                          deposito_lon=EXCLUDED.deposito_lon, activo=EXCLUDED.activo, dummy_prefijo=EXCLUDED.dummy_prefijo""",
                                (co, cli, vt, datos.get("vehiculo_dummy_sr") or None,
                                 _txt(datos.get("vehiculo_dummy_nombre")) or None,
                                 _txt(datos.get("deposito_nombre")) or None,
                                 datos.get("deposito_lat"), datos.get("deposito_lon"),
                                 bool(datos.get("activo", True)), (_txt(datos.get("dummy_prefijo")).upper() or None)))
                    if vt:
                        # Así la importación de rutas en Despacho reconoce a qué cliente pertenece el visit_type.
                        cur.execute("INSERT INTO cat_mapeo_cliente_sr (visit_type, cliente) VALUES (%s,%s) "
                                    "ON CONFLICT (visit_type) DO NOTHING", (vt, cli))
                    _bitacora(cur, usuario, "config_cliente", detalle=f"{co} -> {cli}")
                conn.commit()
            except Exception:
                conn.rollback()
                raise
        return True, "Configuración guardada."
    except Exception as e:
        return False, f"No se pudo guardar la configuración ({type(e).__name__})."


# ---------------------------------------------------------------------------
# Lectura del Excel (plantilla del portal de pedidos)
# ---------------------------------------------------------------------------
CAMPOS_EXCEL = {
    "codigo_pedido": "Código de Pedido",
    "ref": "Ref .adicional de Despacho",
    "cliente_orden": "Cliente Orden",
    "duracion": "Duración (HH:mm:ss)",
    "placa": "Placa",
    "conductor": "Conductor",
    "grupo": "G. Ruta",
    "obs": "Observaciones",
    "fecha": "Fecha de Entrega (dd/mm/aaaa)",
    "ini": "Ini. Horario Entrega (hh:mm)",
    "fin": "Fin Horario Entrega (hh:mm)",
    "cod": "Cód. Domicilio Entrega",
    "nombre": "Nombre Domicilio Entrega",
    "direccion": "Dirección Domicilio Entrega",
    "municipio": "Municipio Entrega",
    "depto": "Departamento Entrega",
    "lat": "Latitud Entrega",
    "lon": "Longitud Entrega",
    "dep_salida": "Depósito Salida",
    "dep_llegada": "Depósito Llegada",
}
CAMPOS_OBLIGATORIOS_EXCEL = ("cliente_orden", "fecha", "cod")


def _norm_encabezado(s):
    return re.sub(r"[^a-z0-9]", "", _norm(s))


def leer_excel_plan(contenido):
    """Lee la hoja PEDIDOS. Devuelve (filas_crudas, errores_globales). No valida contra la base:
    solo extrae los valores de las columnas que usa Control de Ruta (el resto se ignora)."""
    try:
        import openpyxl
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        return [], ["No pude abrir el archivo. Debe ser un Excel (.xlsx) con la plantilla de carga masiva."]
    filas, errores = [], []
    try:
        ws = wb["PEDIDOS"] if "PEDIDOS" in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        try:
            cab = next(it)
        except StopIteration:
            return [], ["La hoja está vacía."]
        mapa_enc = {_norm_encabezado(v): i for i, v in enumerate(cab) if v is not None}
        idx = {}
        for clave, titulo in CAMPOS_EXCEL.items():
            i = mapa_enc.get(_norm_encabezado(titulo))
            if i is not None:
                idx[clave] = i
        faltan = [CAMPOS_EXCEL[c] for c in CAMPOS_OBLIGATORIOS_EXCEL if c not in idx]
        if faltan:
            return [], ["No encuentro estas columnas en la fila 1: " + ", ".join(faltan) +
                        ". ¿Es la plantilla de carga masiva?"]
        for n, fila in enumerate(it, start=2):
            if n - 1 > MAX_FILAS_EXCEL:
                errores.append(f"El archivo tiene más de {MAX_FILAS_EXCEL} filas; se leyeron solo las primeras {MAX_FILAS_EXCEL}.")
                break
            raw = {k: (fila[i] if i < len(fila) else None) for k, i in idx.items()}
            if not any(_txt(raw.get(k)) for k in ("cliente_orden", "fecha", "cod", "nombre")):
                continue   # fila vacía
            raw["n"] = n
            filas.append(raw)
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return filas, errores


def _parse_fecha(v):
    if v is None or _txt(v) == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _txt(v)
    for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(s, f).date()
        except ValueError:
            pass
    return None


def _parse_hora(v):
    """Devuelve 'HH:MM:SS', '' si viene vacío, o None si es inválida."""
    if v is None or _txt(v) == "":
        return ""
    if isinstance(v, datetime):
        return v.strftime("%H:%M:%S")
    if isinstance(v, dtime):
        return v.strftime("%H:%M:%S")
    if isinstance(v, timedelta):
        t = int(v.total_seconds())
        return f"{t // 3600:02d}:{(t % 3600) // 60:02d}:{t % 60:02d}"
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 0 <= v < 1:
        t = int(round(v * 86400))
        return f"{t // 3600:02d}:{(t % 3600) // 60:02d}:{t % 60:02d}"
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?$", _txt(v))
    if not m:
        return None
    h, mi, se = int(m.group(1)), int(m.group(2)), int(m.group(3) or 0)
    if h > 23 or mi > 59 or se > 59:
        return None
    return f"{h:02d}:{mi:02d}:{se:02d}"


def _parse_entero(v):
    """G. Ruta: vacío -> (None, ok); entero (o 3.0 / '3') -> (3, ok); otro -> (None, error)."""
    s = _txt(v)
    if s == "":
        return None, True
    try:
        f = float(s.replace(",", "."))
    except ValueError:
        return None, False
    if not f.is_integer() or f < 0:
        return None, False
    return int(f), True


def _parse_float(v):
    s = _txt(v)
    if s == "":
        return None, True
    try:
        return float(s.replace(",", ".")), True
    except ValueError:
        return None, False


def _buscar_destino(tiendas, cod, nombre):
    """tiendas: lista de dicts (tienda, codigo, lat, lon, direccion) del cliente. La IDENTIDAD de un destino es cliente + código:
    con código numérico solo se busca por código (el mismo código puede existir en otro cliente sin problema).
    Devuelve (tienda|None, conflicto|None). Si el código no existe pero el NOMBRE sí: si esa tienda no tenía código se le adopta;
    si ya tenía otro código es un conflicto (no se adivina cuál es el bueno)."""
    cod = _txt(cod)
    nom_n = _norm(nombre)
    if cod.isdigit():
        n = int(cod)
        for t in tiendas:
            if t["codigo"] == n:
                return t, None
        for t in tiendas:
            if nom_n and _norm(t["tienda"]) == nom_n:
                if t["codigo"] is None:
                    return dict(t, adopta_codigo=n), None
                return None, (f"Ya existe una tienda llamada '{t['tienda']}' con el código {t['codigo']}, pero el Excel trae el código {n}. "
                              "El código es lo que identifica a la tienda: revisa cuál es el correcto.")
        return None, None
    # Código no numérico (poco común): no hay código que comparar, se identifica por nombre o por el propio texto.
    cod_n = _norm(cod)
    for t in tiendas:
        tn = _norm(t["tienda"])
        if (nom_n and tn == nom_n) or (cod_n and tn == cod_n):
            return t, None
    return None, None


def _clave_destino(codigo, destino):
    return str(codigo) if codigo not in (None, "") else _norm(destino)


# ---------------------------------------------------------------------------
# Validación (devuelve las filas con errores/avisos; NO guarda nada)
# ---------------------------------------------------------------------------
def validar_filas(get_conn_fn, filas, prueba=False, permitir_repetidas=False, hoy=None, clientes_permitidos=None):
    hoy = hoy or hoy_guatemala()
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT cliente_orden, cliente, vehiculo_dummy_sr, deposito_nombre, deposito_lat, deposito_lon, "
                    "visit_type FROM cat_transporte_cliente WHERE activo")
        cfg = {}
        for r in cur.fetchall():
            cfg[_clave_cliente(r[0])] = dict(cliente_orden=r[0], cliente=r[1], dummy=r[2], dep=r[3],
                                             dep_lat=r[4], dep_lon=r[5], visit_type=r[6], auto=False)
        # Clientes no configurados: si se pueden identificar sin duda por su nombre, se usan igual (se guardan al guardar el plan).
        for raw_ in filas:
            co_ = _txt(raw_.get("cliente_orden"))
            if co_ and _clave_cliente(co_) not in cfg:
                cli_ = _cliente_cr_por_nombre(cur, co_)
                if cli_:
                    cfg[_clave_cliente(co_)] = dict(cliente_orden=co_, cliente=cli_, dummy=None, dep=None, dep_lat=None, dep_lon=None,
                                                    visit_type=_prefijo_codigo(co_).lower().replace(" ", ""), auto=True)
        clientes_cr = sorted({c["cliente"] for c in cfg.values()})
        tiendas = {}
        if clientes_cr:
            cur.execute("SELECT cliente, tienda, codigo_tienda, latitud, longitud, direccion FROM cat_clientes_tiendas "
                        "WHERE cliente = ANY(%s)", (clientes_cr,))
            for cli, tienda, codigo, la, lo, di in cur.fetchall():
                tiendas.setdefault(cli, []).append(dict(tienda=tienda, codigo=codigo, lat=la, lon=lo, direccion=di))
        cur.execute("SELECT placa, piloto, id_sr, sincronizado_sr FROM cat_camiones")
        camiones = {sr_int.clave_placa(p): dict(placa=p, piloto=pi, id_sr=i, sync=s) for p, pi, i, s in cur.fetchall()}
        cur.execute("SELECT nombre, id_sr FROM cat_pilotos")
        pilotos = {_norm(n): dict(nombre=n, id_sr=i) for n, i in cur.fetchall()}

        # Para detectar duplicados contra lo ya guardado y grupos ya despachados.
        claves_plan = set()
        out = []
        sin_pool_cache = {}
        cur.execute("SELECT cliente_orden, dummy_prefijo FROM cat_transporte_cliente WHERE dummy_prefijo IS NOT NULL")
        dummy_pref = dict(cur.fetchall())
        for raw in filas:
            f = dict(n=raw.get("n"), errores=[], avisos=[])
            # --- cliente ---
            co_raw = _txt(raw.get("cliente_orden"))
            c = cfg.get(_clave_cliente(co_raw)) if co_raw else None
            if not co_raw:
                f["errores"].append("Falta el Cliente Orden.")
            elif not c:
                f["errores"].append(f"El Cliente Orden '{co_raw}' no está configurado para Transporte (Configuración).")
            if c and clientes_permitidos is not None and c["cliente"] not in set(clientes_permitidos):
                f["errores"].append(f"No tienes acceso al cliente '{c['cliente']}' (Cliente Orden {c['cliente_orden']}).")
                c = None
            f["cliente_orden"] = c["cliente_orden"] if c else co_raw
            f["cliente"] = c["cliente"] if c else None
            # --- fecha ---
            fecha = _parse_fecha(raw.get("fecha"))
            if fecha is None:
                f["errores"].append("La fecha de entrega falta o no es válida (dd/mm/aaaa).")
            elif fecha < hoy:
                f["errores"].append(f"La fecha {fecha:%d/%m/%Y} ya pasó.")
            f["fecha"] = fecha
            # --- grupo ---
            grupo, ok_g = _parse_entero(raw.get("grupo"))
            if not ok_g:
                f["errores"].append(f"G. Ruta debe ser un número entero (vino '{_txt(raw.get('grupo'))}').")
            f["grupo"] = grupo
            # --- destino ---
            cod, nombre = _txt(raw.get("cod")), _txt(raw.get("nombre"))
            if not cod:
                f["errores"].append("Falta el código del domicilio de entrega.")
            lat, ok_la = _parse_float(raw.get("lat"))
            lon, ok_lo = _parse_float(raw.get("lon"))
            if not ok_la or not ok_lo:
                f["errores"].append("La latitud o la longitud no son números.")
                lat = lon = None
            elif (lat is None) != (lon is None):
                f["errores"].append("Falta la latitud o la longitud (deben venir las dos).")
                lat = lon = None
            elif lat is not None and not (-90 <= lat <= 90 and -180 <= lon <= 180):
                f["errores"].append("Las coordenadas están fuera de rango.")
                lat = lon = None
            elif lat is not None and not (13.5 <= lat <= 18.5 and -92.5 <= lon <= -88.0):
                f["avisos"].append("Las coordenadas quedan fuera de Guatemala; revisa que no estén invertidas.")
            dest = None
            conflicto_dest = None
            if c and cod:
                dest, conflicto_dest = _buscar_destino(tiendas.get(c["cliente"], []), cod, nombre)
            if conflicto_dest:
                f["errores"].append(conflicto_dest)
            f["dest_nuevo"] = False
            f["dest_codigo"] = int(cod) if cod.isdigit() else None
            if dest:
                f["destino"] = dest["tienda"]
                if dest.get("adopta_codigo") is not None:
                    f["dest_adopta_codigo"] = True
                    f["avisos"].append(f"La tienda '{dest['tienda']}' no tenía código: se le asignará el {dest['adopta_codigo']}.")
                elif nombre and _norm(nombre) != _norm(dest["tienda"]):
                    f["avisos"].append(f"El código {cod} ya existe como '{dest['tienda']}' (el Excel dice '{nombre}'): se usa el nombre del catálogo.")
                f["lat"] = dest["lat"] if dest["lat"] is not None else lat
                f["lon"] = dest["lon"] if dest["lon"] is not None else lon
                f["dest_actualiza_coord"] = dest["lat"] is None and lat is not None
                if f["lat"] is None:
                    f["coord_pendiente_sr"] = True
                    f["avisos"].append("Sin coordenadas: se tomarán de SimpliRoute (visitas de los últimos días) al enviar el plan.")
                if dest["codigo"] is not None:
                    f["dest_codigo"] = dest["codigo"]
            elif c and cod and not conflicto_dest:
                f["dest_nuevo"] = True
                f["destino"] = nombre or cod
                f["lat"], f["lon"] = lat, lon
                if not nombre:
                    f["errores"].append("Destino nuevo: falta el nombre del domicilio.")
                if lat is None:
                    f["coord_pendiente_sr"] = True
                    f["avisos"].append("Destino nuevo sin coordenadas: se agregará al catálogo y sus coordenadas se tomarán de SimpliRoute al enviar.")
                else:
                    f["avisos"].append("Destino nuevo: se agregará al catálogo de tiendas del cliente.")
            else:
                f["destino"], f["lat"], f["lon"] = nombre or cod, lat, lon
            f["cod"] = cod
            partes = [_txt(raw.get(k)) for k in ("direccion", "municipio", "depto")]
            f["direccion_excel"] = ", ".join(p for p in partes if p)
            f["direccion"] = f["direccion_excel"] or ((f["destino"] + ", Guatemala") if f["destino"] else "")
            if dest and f["direccion_excel"] and not (dest.get("direccion") or "").strip():
                f["dest_actualiza_dir"] = True                 # la tienda existe pero no tiene dirección: se completa
            # --- horas y duración ---
            for k, etiqueta in (("ini", "El inicio del horario"), ("fin", "El fin del horario"), ("duracion", "La duración")):
                h = _parse_hora(raw.get(k))
                if h is None:
                    f["errores"].append(f"{etiqueta} no tiene formato hh:mm.")
                    h = ""
                f[k] = h
            f["obs"] = _txt(raw.get("obs"))
            f["ref"] = _txt(raw.get("ref"))
            # --- vehículo del viaje (columna Placa): genérico (DC01, ADICIONALES1...) o camión real ---
            placa_raw = _txt(raw.get("placa"))
            f["placa"] = None            # camión REAL asignado (None si es genérico)
            f["vehiculo"] = None         # lo que trae la columna, ya normalizado
            f["vehiculo_sr_id"] = None
            f["generico"] = False
            if placa_raw:
                cam = camiones.get(sr_int.clave_placa(placa_raw))
                if not cam:
                    f["errores"].append(f"El vehículo '{placa_raw}' no está en el catálogo de Control de Ruta. "
                                        "Si es un genérico de SimpliRoute, sincroniza camiones (Catálogos → Integración SimpliRoute).")
                elif not cam["id_sr"]:
                    f["errores"].append(f"El vehículo {cam['placa']} no está sincronizado con SimpliRoute; sincroniza camiones primero.")
                else:
                    f["vehiculo"] = cam["placa"]
                    f["vehiculo_sr_id"] = int(cam["id_sr"])
                    f["generico"] = es_generico(cam["placa"])
                    if not f["generico"]:
                        f["placa"] = cam["placa"]
            elif c:
                if c["cliente_orden"] not in sin_pool_cache:
                    sin_pool_cache[c["cliente_orden"]] = not (c.get("auto") is False and pool_propio(get_conn_fn, {
                        "dummy_prefijo": dummy_pref.get(c["cliente_orden"]), "vehiculo_dummy_sr": c.get("dummy"), "vehiculo_dummy_nombre": ""}))
                if sin_pool_cache[c["cliente_orden"]]:
                    f["errores"].append("Falta el vehículo en la columna Placa (error de plantilla): usa el genérico, como DC01 o "
                                        "ADICIONALES1, o el camión real.")
            cond_raw = _txt(raw.get("conductor"))
            f["conductor"] = None
            if cond_raw and _norm(cond_raw) not in ("sin asignar", "sinasignar"):
                p = pilotos.get(_norm(cond_raw))
                if not p:
                    f["errores"].append(f"El conductor '{cond_raw}' no existe en el catálogo de pilotos.")
                else:
                    f["conductor"] = p["nombre"]
            f["dep_salida"] = _txt(raw.get("dep_salida"))
            f["dep_llegada"] = _txt(raw.get("dep_llegada"))
            f["_cfg"] = c
            if c and fecha:
                claves_plan.add((c["cliente_orden"], fecha))
            out.append(f)

        # --- consistencia dentro de cada grupo (mismo cliente + fecha + G. Ruta) ---
        grupos = {}
        for f in out:
            if f["grupo"] is not None and f["fecha"] and f["cliente"]:
                grupos.setdefault((f["cliente_orden"], f["fecha"], f["grupo"]), []).append(f)
        for (co, fe, gr), lista in grupos.items():
            for campo, etiqueta in (("vehiculo", "vehículos (Placa)"), ("conductor", "conductores"), ("dep_salida", "depósitos de salida")):
                vals = {f[campo] for f in lista if f[campo]}
                if len(vals) > 1:
                    for f in lista:
                        f["errores"].append(f"El grupo {gr} trae {etiqueta} distintos: {', '.join(sorted(vals))}.")

        # --- un camión real en demasiados viajes el mismo día (casi seguro: placa genérica puesta como real, o error de plantilla) ---
        viajes_por_placa = {}
        for f in out:
            if f.get("placa") and f["fecha"] and f["cliente"]:
                viajes_por_placa.setdefault((f["placa"], f["fecha"]), set()).add(
                    ("g", f["cliente_orden"], f["grupo"]) if f["grupo"] is not None else ("fila", f["n"]))
        if viajes_por_placa:
            cur.execute("""SELECT v.placa, p.fecha, COUNT(*) FROM plan_viaje v JOIN plan_transporte p ON p.id = v.plan_id
                           WHERE v.placa = ANY(%s) AND p.fecha = ANY(%s) AND v.estado <> 'cancelado' AND p.prueba = %s GROUP BY v.placa, p.fecha""",
                        (sorted({k[0] for k in viajes_por_placa}), sorted({k[1] for k in viajes_por_placa}), bool(prueba)))
            ya_en_bd = {(a, b): c_ for a, b, c_ in cur.fetchall()}
            for (placa_k, fecha_k), viajes_k in viajes_por_placa.items():
                total_k = len(viajes_k) + ya_en_bd.get((placa_k, fecha_k), 0)
                if total_k > MAX_VIAJES_POR_PLACA_DIA:
                    for f in out:
                        if f.get("placa") == placa_k and f["fecha"] == fecha_k:
                            f["errores"].append(
                                f"El camión {placa_k} quedaría en {total_k} viajes el {fecha_k:%d/%m/%Y} (máximo {MAX_VIAJES_POR_PLACA_DIA}). "
                                "Si es un vehículo genérico (como DC01 o ADICIONALES1) ponlo con su nombre; si es un camión real, "
                                "cada viaje debe llevar el suyo o dejarse sin asignar.")

        # --- duplicados dentro del archivo ---
        if not permitir_repetidas:
            vistos = {}
            for f in out:
                if f["cliente"] and f["fecha"] and f["cod"]:
                    k = (f["cliente_orden"], f["fecha"], f["grupo"], _clave_destino(f["dest_codigo"], f["destino"]))
                    if k in vistos:
                        f["errores"].append(f"Repetida en el archivo (igual a la fila {vistos[k]}). "
                                            "Si es correcto, activa 'Permitir paradas repetidas'.")
                    else:
                        vistos[k] = f["n"]

        # --- contra lo ya guardado en la base ---
        if claves_plan:
            cos = sorted({k[0] for k in claves_plan})
            fechas = sorted({k[1] for k in claves_plan})
            cur.execute("""
                SELECT p.cliente_orden, p.fecha, v.grupo, v.estado, pa.cod_destino, pa.destino, pa.estado
                FROM plan_parada pa
                JOIN plan_viaje v ON v.id = pa.viaje_id
                JOIN plan_transporte p ON p.id = v.plan_id
                WHERE p.prueba = %s AND p.cliente_orden = ANY(%s) AND p.fecha = ANY(%s)""",
                        (bool(prueba), cos, fechas))
            existentes = set()
            cerrados = set()
            abiertos = set()
            for co, fe, gr, est_v, cdest, dest, est_p in cur.fetchall():
                if gr is not None:
                    (cerrados if est_v == "despachado" else abiertos if est_v != "cancelado" else set()).add((co, fe, gr))
                if est_p != "cancelada" and est_v != "cancelado":
                    existentes.add((co, fe, gr, _clave_destino(cdest, dest)))
            for f in out:
                if not (f["cliente"] and f["fecha"]):
                    continue
                if f["grupo"] is not None:
                    kg = (f["cliente_orden"], f["fecha"], f["grupo"])
                    if kg in cerrados and kg not in abiertos:
                        f["errores"].append(f"El viaje del grupo {f['grupo']} ya fue despachado; no se le pueden agregar paradas.")
                if not permitir_repetidas and f["cod"]:
                    k = (f["cliente_orden"], f["fecha"], f["grupo"], _clave_destino(f["dest_codigo"], f["destino"]))
                    if k in existentes:
                        f["errores"].append("Ya está cargada en este plan. Si es un viaje adicional al mismo destino, "
                                            "activa 'Permitir paradas repetidas'.")
    for f in out:
        f["ok"] = not f["errores"]
    return out


def agrupar_para_vista(filas_ok):
    """Resumen para la vista previa: {(cliente_orden, fecha): {'viajes': n, 'paradas': m}}."""
    res = {}
    vistos = {}
    for f in filas_ok:
        k = (f["cliente_orden"], f["fecha"])
        d = res.setdefault(k, {"viajes": 0, "paradas": 0})
        d["paradas"] += 1
        if f["grupo"] is None:
            d["viajes"] += 1
        else:
            if (k, f["grupo"]) not in vistos:
                vistos[(k, f["grupo"])] = True
                d["viajes"] += 1
    return res


# ---------------------------------------------------------------------------
# Guardado como borrador (antes de tocar SR)
# ---------------------------------------------------------------------------
def guardar_filas(get_conn_fn, filas_ok, usuario, prueba=False, archivo=None):
    """Guarda en Supabase las filas YA validadas. Todo en una transacción por plan, con candado para que dos
    coordinadores a la vez no repitan un consecutivo. Devuelve dict {ok, planes: [...], mensaje}."""
    if not filas_ok:
        return {"ok": False, "planes": [], "mensaje": "No hay filas válidas para guardar."}
    por_plan = {}
    for f in sorted(filas_ok, key=lambda x: (x["n"] or 0)):
        por_plan.setdefault((f["cliente_orden"], f["fecha"]), []).append(f)
    resultado = []
    try:
        with closing(get_conn_fn()) as conn:
            try:
                with conn.cursor() as cur:
                    cur.execute("INSERT INTO plan_carga (usuario, archivo, prueba, filas) VALUES (%s,%s,%s,%s) RETURNING id",
                                (usuario, (archivo or "(captura manual)")[:200], bool(prueba), len(filas_ok)))
                    carga_id = cur.fetchone()[0]
                    for (co, fecha), lista in por_plan.items():
                        cfg = lista[0]["_cfg"]
                        if cfg.get("auto"):
                            cur.execute("""INSERT INTO cat_transporte_cliente (cliente_orden, cliente, visit_type, activo)
                                           VALUES (%s,%s,%s,TRUE) ON CONFLICT (cliente_orden) DO NOTHING""",
                                        (cfg["cliente_orden"], cfg["cliente"], cfg["visit_type"]))
                            cur.execute("INSERT INTO cat_mapeo_cliente_sr (visit_type, cliente) VALUES (%s,%s) ON CONFLICT (visit_type) DO NOTHING",
                                        (cfg["visit_type"], cfg["cliente"]))
                            _bitacora(cur, usuario, "cliente_automatico", detalle=f"{cfg['cliente_orden']} -> {cfg['cliente']}")
                        nom = nombre_plan(co, fecha, prueba)
                        cur.execute("SELECT pg_advisory_xact_lock(hashtext(%s))", (nom,))
                        cur.execute("""INSERT INTO plan_transporte (cliente_orden, cliente, fecha, nombre, prueba, creado_por)
                                       VALUES (%s,%s,%s,%s,%s,%s) ON CONFLICT (nombre) DO NOTHING""",
                                    (co, cfg["cliente"], fecha, nom, bool(prueba), usuario))
                        cur.execute("SELECT id FROM plan_transporte WHERE nombre = %s", (nom,))
                        plan_id = cur.fetchone()[0]
                        cur.execute("SELECT COALESCE(MAX(num),0) FROM plan_viaje WHERE plan_id = %s", (plan_id,))
                        num_max = cur.fetchone()[0]
                        cur.execute("SELECT COALESCE(MAX(consecutivo),0) FROM plan_parada WHERE plan_id = %s", (plan_id,))
                        consec = cur.fetchone()[0]
                        viajes_nuevos = paradas_nuevas = paradas_a_existente = 0
                        viaje_de_grupo = {}
                        cur.execute("SELECT grupo, id FROM plan_viaje WHERE plan_id = %s AND grupo IS NOT NULL "
                                    "AND estado NOT IN ('cancelado','despachado')", (plan_id,))
                        for g, vid in cur.fetchall():
                            viaje_de_grupo[g] = vid
                        for f in lista:
                            if f.get("dest_nuevo") or f.get("dest_actualiza_coord") or f.get("dest_actualiza_dir") or f.get("dest_adopta_codigo"):
                                _upsert_tienda(cur, cfg["cliente"], f)
                            vid = viaje_de_grupo.get(f["grupo"]) if f["grupo"] is not None else None
                            if vid is None:
                                num_max += 1
                                dep = f["dep_salida"] or cfg["dep"]
                                cur.execute("""
                                    INSERT INTO plan_viaje (plan_id, num, grupo, deposito_nombre, deposito_lat, deposito_lon)
                                    VALUES (%s,%s,%s,%s,%s,%s) RETURNING id""",
                                            (plan_id, num_max, f["grupo"], dep, cfg["dep_lat"], cfg["dep_lon"]))
                                vid = cur.fetchone()[0]
                                viajes_nuevos += 1
                                if f["grupo"] is not None:
                                    viaje_de_grupo[f["grupo"]] = vid
                                orden = 1
                            else:
                                cur.execute("SELECT COALESCE(MAX(orden),0)+1 FROM plan_parada WHERE viaje_id = %s", (vid,))
                                orden = cur.fetchone()[0]
                                paradas_a_existente += 1
                            _aplicar_vehiculo_fila(cur, vid, f, usuario)
                            consec += 1
                            cur.execute("""
                                INSERT INTO plan_parada (viaje_id, plan_id, orden, consecutivo, codigo_pedido, cod_destino,
                                    destino, direccion, latitud, longitud, ventana_ini, ventana_fin, duracion,
                                    observaciones, ref_adicional, carga_id)
                                VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
                                        (vid, plan_id, orden, consec, codigo_pedido(co, fecha, consec),
                                         str(f["dest_codigo"]) if f["dest_codigo"] is not None else (f["cod"] or None),
                                         f["destino"], f["direccion"], f["lat"], f["lon"],
                                         f["ini"] or None, f["fin"] or None, f["duracion"] or None,
                                         f["obs"] or None, f["ref"] or None, carga_id))
                            paradas_nuevas += 1
                        _bitacora(cur, usuario, "cargar_plan", plan_id=plan_id,
                                  detalle=f"{paradas_nuevas} parada(s), {viajes_nuevos} viaje(s) nuevo(s), "
                                          f"{paradas_a_existente} agregada(s) a un viaje existente")
                        resultado.append({"plan_id": plan_id, "nombre": nom, "viajes_nuevos": viajes_nuevos,
                                          "paradas": paradas_nuevas, "a_viaje_existente": paradas_a_existente})
                conn.commit()
            except Exception:
                conn.rollback()
                raise
    except Exception as e:
        return {"ok": False, "planes": [], "mensaje": f"No se pudo guardar el plan ({type(e).__name__}). No se guardó nada.",
                "excepcion": e}
    return {"ok": True, "planes": resultado, "carga_id": carga_id, "mensaje": "Plan guardado como borrador."}


def _aplicar_vehiculo_fila(cur, viaje_id, f, usuario):
    """Vehículo de la fila (columna Placa) en el viaje: un genérico queda como el vehículo con el que se crea la ruta en SR (el viaje
    sigue «sin camión»); un camión real queda asignado. Nunca pisa lo que el viaje ya tenga."""
    cur.execute("SELECT placa, dummy_sr_id, piloto FROM plan_viaje WHERE id = %s", (viaje_id,))
    placa_act, dummy_act, piloto_act = cur.fetchone()
    if f.get("vehiculo_sr_id"):
        if f["generico"]:
            if dummy_act is None and placa_act is None:
                cur.execute("UPDATE plan_viaje SET dummy_sr_id = %s WHERE id = %s", (f["vehiculo_sr_id"], viaje_id))
        elif placa_act is None:
            piloto = f["conductor"]
            if not piloto:                       # sin conductor en el Excel: el piloto habitual de ese camión en Control de Ruta
                cur.execute("SELECT piloto FROM cat_camiones WHERE placa = %s", (f["placa"],))
                r_ = cur.fetchone()
                if r_ and r_[0]:
                    cur.execute("SELECT 1 FROM cat_pilotos WHERE nombre = %s", (r_[0],))
                    piloto = r_[0] if cur.fetchone() else None
            _fijar_camion_piloto(cur, viaje_id, f["placa"], piloto, usuario)
            return
    if f.get("conductor") and piloto_act is None:
        cur.execute("SELECT id_sr FROM cat_pilotos WHERE nombre = %s", (f["conductor"],))
        r = cur.fetchone()
        cur.execute("UPDATE plan_viaje SET piloto = %s, piloto_sr_id = %s WHERE id = %s", (f["conductor"], r[0] if r else None, viaje_id))


def _upsert_tienda(cur, cliente, f):
    """Destino nuevo (o al que le falta algo): se agrega/completa en el catálogo. Se identifica por cliente + CÓDIGO (el nombre solo
    se usa si no hay código). El km no se inventa: una tienda nueva entra con el km vacío."""
    cod = f["dest_codigo"]
    fila = None
    if cod is not None:
        cur.execute("SELECT tienda FROM cat_clientes_tiendas WHERE cliente = %s AND codigo_tienda = %s", (cliente, cod))
        fila = cur.fetchone()
    if fila is None:
        cur.execute("SELECT tienda FROM cat_clientes_tiendas WHERE cliente = %s AND tienda = %s", (cliente, f["destino"]))
        fila = cur.fetchone()
    if fila:
        cur.execute("UPDATE cat_clientes_tiendas SET latitud = COALESCE(latitud, %s), longitud = COALESCE(longitud, %s), "
                    "direccion = COALESCE(NULLIF(TRIM(direccion), ''), %s), codigo_tienda = COALESCE(codigo_tienda, %s) "
                    "WHERE cliente = %s AND tienda = %s",
                    (f["lat"], f["lon"], f.get("direccion_excel") or None, cod, cliente, fila[0]))
    else:
        cur.execute("INSERT INTO cat_clientes_tiendas (cliente, tienda, km, codigo_tienda, clasificacion, latitud, longitud, direccion) "
                    "VALUES (%s,%s,NULL,%s,'Local',%s,%s,%s)",
                    (cliente, f["destino"], cod, f["lat"], f["lon"], f.get("direccion_excel") or None))


def validar_coordenadas_catalogo(df):
    """Carga masiva de «Clientes y Tiendas»: valida las columnas OPCIONALES latitud/longitud antes de guardar.
    Devuelve (df_limpio, errores, avisos). Celdas vacías = «sin dato» (no borran lo que ya hay en el catálogo).
    Errores (bloquean la carga): una sola de las dos, texto, fuera de rango, 0/0, o latitud y longitud invertidas.
    Avisos (no bloquean): coordenadas que caen fuera de Guatemala."""
    if "latitud" not in df.columns and "longitud" not in df.columns:
        return df, [], []
    df = df.copy()
    errores, avisos = [], []
    if "latitud" not in df.columns or "longitud" not in df.columns:
        return df, ["El archivo trae solo una de las dos columnas: latitud y longitud deben venir juntas."], []
    lats, lons = [], []
    for i, fila in df.iterrows():
        n = int(i) + 2                                           # fila en el Excel (la 1 es el encabezado)
        tienda = _txt(fila.get("tienda")) or "(sin nombre)"
        et = f"Fila {n} ({tienda})"
        la, ok_la = _parse_float(fila.get("latitud"))
        lo, ok_lo = _parse_float(fila.get("longitud"))
        if not ok_la or not ok_lo:
            errores.append(f"{et}: latitud o longitud no son números.")
            la = lo = None
        elif (la is None) != (lo is None):
            errores.append(f"{et}: trae solo una de las dos (latitud y longitud van juntas).")
            la = lo = None
        elif la is not None:
            if not (-90 <= la <= 90 and -180 <= lo <= 180):
                errores.append(f"{et}: coordenadas fuera de rango ({la}, {lo}).")
                la = lo = None
            elif la == 0 and lo == 0:
                errores.append(f"{et}: las coordenadas son 0, 0 (celda sin llenar).")
                la = lo = None
            elif abs(la) > 40 and abs(lo) < 40:
                errores.append(f"{et}: latitud y longitud parecen estar invertidas ({la}, {lo}).")
                la = lo = None
            else:
                if not (13.5 <= la <= 18.5 and -92.5 <= lo <= -88.0):
                    avisos.append(f"{et}: queda fuera de Guatemala ({la}, {lo}); revisa que sea correcto.")
                la, lo = round(la, 6), round(lo, 6)
        lats.append(la)
        lons.append(lo)
    df["latitud"] = lats
    df["longitud"] = lons
    return df, errores, avisos


def codigos_repetidos(get_conn_fn):
    """Códigos que tienen más de una tienda dentro del mismo cliente: [(cliente, codigo, [tiendas])]."""
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT cliente, codigo_tienda, ARRAY_AGG(tienda ORDER BY tienda) FROM cat_clientes_tiendas "
                    "WHERE codigo_tienda IS NOT NULL GROUP BY cliente, codigo_tienda HAVING COUNT(*) > 1 ORDER BY cliente, codigo_tienda")
        return [(a, b, list(c)) for a, b, c in cur.fetchall()]


def resolver_carga_tiendas(get_conn_fn, df):
    """Carga masiva de «Clientes y Tiendas»: la tienda se identifica por CLIENTE + CÓDIGO. Antes de guardar:
    - si el código ya existe en ese cliente, la fila actualiza ESA tienda (se conserva el nombre del catálogo) aunque el nombre venga distinto;
    - si el nombre ya existe con OTRO código, o un código se repite en el archivo, es un error (no se adivina);
    - sin código, se identifica por nombre (y se avisa).
    Devuelve (df_resuelto, errores, avisos, resumen)."""
    df = df.copy()
    errores, avisos = [], []
    resumen = {"nuevas": 0, "existentes": 0, "sin_codigo": 0, "nombre_distinto": 0}
    if df.empty or "cliente" not in df.columns or "tienda" not in df.columns:
        return df, errores, avisos, resumen
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT nombre FROM cat_clientes")
        cliente_real = {_norm(r[0]): r[0] for r in cur.fetchall()}
        cur.execute("SELECT cliente, tienda, codigo_tienda FROM cat_clientes_tiendas")
        por_codigo, por_nombre = {}, {}
        for cli, tienda, cod in cur.fetchall():
            if cod is not None:
                por_codigo[(_norm(cli), int(cod))] = tienda
            por_nombre[(_norm(cli), _norm(tienda))] = (tienda, cod)
    vistos_codigo, vistos_nombre = {}, {}
    nuevos_nombres = []
    for i, fila in df.iterrows():
        n = int(i) + 2
        cli_n, tienda = _norm(fila.get("cliente")), _txt(fila.get("tienda"))
        et = f"Fila {n} ({tienda or 'sin nombre'})"
        if not tienda:
            errores.append(f"{et}: falta el nombre de la tienda.")
            nuevos_nombres.append(tienda)
            continue
        cod_txt = _txt(fila.get("codigo_tienda"))
        cod = None
        if cod_txt:
            try:
                cod = int(float(cod_txt))
            except ValueError:
                errores.append(f"{et}: el código '{cod_txt}' no es un número.")
        nuevo_nombre = tienda
        if cod is not None:
            k = (cli_n, cod)
            if k in vistos_codigo:
                errores.append(f"{et}: el código {cod} se repite en el archivo (también en la fila {vistos_codigo[k]}).")
            vistos_codigo[k] = n
            if k in por_codigo:
                nuevo_nombre = por_codigo[k]
                resumen["existentes"] += 1
                if _norm(nuevo_nombre) != _norm(tienda):
                    resumen["nombre_distinto"] += 1
                    avisos.append(f"{et}: el código {cod} ya es '{nuevo_nombre}'; se actualiza esa tienda y se conserva su nombre.")
            else:
                ya = por_nombre.get((cli_n, _norm(tienda)))
                if ya and ya[1] is not None and int(ya[1]) != cod:
                    errores.append(f"{et}: ya existe con el código {ya[1]} y el archivo trae el {cod}. Revisa cuál es el correcto.")
                elif ya:
                    nuevo_nombre = ya[0]
                    resumen["existentes"] += 1                       # misma tienda sin código: se le asigna este
                else:
                    resumen["nuevas"] += 1
        else:
            resumen["sin_codigo"] += 1
            ya = por_nombre.get((cli_n, _norm(tienda)))
            if ya:
                nuevo_nombre = ya[0]
                resumen["existentes"] += 1
            else:
                resumen["nuevas"] += 1
        kn = (cli_n, _norm(nuevo_nombre))
        if kn in vistos_nombre and vistos_nombre[kn] != n:
            errores.append(f"{et}: la tienda aparece dos veces en el archivo (también en la fila {vistos_nombre[kn]}).")
        vistos_nombre[kn] = n
        nuevos_nombres.append(nuevo_nombre)
    df["tienda"] = nuevos_nombres
    if resumen["sin_codigo"]:
        avisos.append(f"{resumen['sin_codigo']} fila(s) vienen SIN código: se identifican por nombre, y escribir una tienda distinta crearía un duplicado. "
                      "Conviene que todas traigan su código.")
    return df, errores, avisos, resumen


def fila_manual(cliente_orden, fecha, cod, nombre, grupo=None, obs="", placa="", conductor=""):
    """Arma una fila 'cruda' como si viniera del Excel, para que la carga manual use la MISMA validación."""
    return {"n": 0, "cliente_orden": cliente_orden, "fecha": fecha, "cod": cod, "nombre": nombre,
            "grupo": grupo if grupo else None, "obs": obs, "placa": placa, "conductor": conductor}


# ---------------------------------------------------------------------------
# Lectura de planes y viajes (para pantalla y reportes)
# ---------------------------------------------------------------------------
def listar_planes(get_conn_fn, clientes_orden=None, fecha_ini=None, fecha_fin=None, prueba=False, limite=60):
    import psycopg2.extras
    where, params = ["p.prueba = %s"], [bool(prueba)]
    if clientes_orden:
        where.append("p.cliente_orden = ANY(%s)")
        params.append(list(clientes_orden))
    if fecha_ini:
        where.append("p.fecha >= %s")
        params.append(fecha_ini)
    if fecha_fin:
        where.append("p.fecha <= %s")
        params.append(fecha_fin)
    with closing(get_conn_fn()) as conn, conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute(f"""
            SELECT p.*,
              COUNT(v.id) FILTER (WHERE v.estado <> 'cancelado') AS viajes,
              COUNT(v.id) FILTER (WHERE v.estado IN ('asignado','despachado')) AS con_camion
            FROM plan_transporte p LEFT JOIN plan_viaje v ON v.plan_id = p.id
            WHERE {' AND '.join(where)}
            GROUP BY p.id ORDER BY p.fecha, p.cliente_orden LIMIT %s""", params + [limite])
        return [dict(r) for r in cur.fetchall()]


def fechas_con_plan(get_conn_fn, clientes_orden, prueba, desde, limite=60):
    """Fechas (desde `desde` en adelante) que tienen al menos un plan de esos clientes. Sirve para no tener que adivinar la fecha."""
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT DISTINCT fecha FROM plan_transporte WHERE prueba = %s AND cliente_orden = ANY(%s) AND fecha >= %s "
                    "ORDER BY fecha LIMIT %s", (bool(prueba), list(clientes_orden), desde, limite))
        return [r[0] for r in cur.fetchall()]


def detalle_plan(get_conn_fn, plan_id):
    """{'plan': {...}, 'viajes': [{..., 'paradas': [...]}]}  (los cancelados al final)."""
    import psycopg2.extras
    with closing(get_conn_fn()) as conn, conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute("SELECT * FROM plan_transporte WHERE id = %s", (plan_id,))
        p = cur.fetchone()
        if not p:
            return None
        cur.execute("SELECT * FROM plan_viaje WHERE plan_id = %s ORDER BY (estado = 'cancelado'), num", (plan_id,))
        viajes = [dict(r) for r in cur.fetchall()]
        cur.execute("SELECT * FROM plan_parada WHERE plan_id = %s ORDER BY viaje_id, orden", (plan_id,))
        paradas = {}
        for r in cur.fetchall():
            paradas.setdefault(r["viaje_id"], []).append(dict(r))
        for v in viajes:
            v["paradas"] = paradas.get(v["id"], [])
    return {"plan": dict(p), "viajes": viajes}


def listar_cargas(get_conn_fn, prueba, limite=10, clientes=None):
    """Últimas plantillas cargadas (más nueva primero): [{id, archivo, momento, usuario, filas}]. Con `clientes`, solo las plantillas que
    tienen paradas de esos clientes (nadie ve plantillas de clientes a los que no tiene acceso)."""
    filtro, params = "", [bool(prueba)]
    if clientes is not None:
        filtro = (" AND EXISTS (SELECT 1 FROM plan_parada pa JOIN plan_transporte p ON p.id = pa.plan_id "
                  "WHERE pa.carga_id = plan_carga.id AND p.cliente = ANY(%s))")
        params.append(list(clientes))
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute(f"SELECT id, archivo, momento, usuario, filas FROM plan_carga WHERE prueba = %s{filtro} ORDER BY id DESC LIMIT %s",
                    params + [limite])
        return [dict(id=a, archivo=b, momento=c, usuario=d, filas=e) for a, b, c, d, e in cur.fetchall()]


def estado_carga(get_conn_fn, carga_id, clientes=None):
    """Qué tan enviada está una plantilla. Un plan cuenta como «enviado» cuando existe en SR y ya no le queda ninguna parada de
    esa plantilla por enviar. Devuelve {planes: [...], planes_total, planes_enviados, paradas_total, paradas_enviadas}."""
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("""
            SELECT p.id, p.nombre, p.cliente, p.cliente_orden, p.fecha, p.plan_id_sr,
                   COUNT(*) FILTER (WHERE pa.estado <> 'cancelada'),
                   COUNT(*) FILTER (WHERE pa.estado = 'en_sr'),
                   COUNT(*) FILTER (WHERE pa.estado IN ('pendiente','enviando','error'))
            FROM plan_parada pa JOIN plan_transporte p ON p.id = pa.plan_id
            WHERE pa.carga_id = %s AND (%s::text[] IS NULL OR p.cliente = ANY(%s))
            GROUP BY p.id ORDER BY p.fecha, p.cliente""", (carga_id, list(clientes) if clientes is not None else None,
                                                          list(clientes) if clientes is not None else None))
        planes = []
        for pid, nombre, cliente, co, fecha, plan_sr, activas, en_sr, pend in cur.fetchall():
            if not activas:
                continue                                          # todo cancelado: no cuenta
            planes.append(dict(id=pid, nombre=nombre, cliente=cliente, cliente_orden=co, fecha=fecha, paradas=activas,
                               paradas_enviadas=en_sr, pendientes=pend, enviado=bool(plan_sr) and pend == 0))
    return {"planes": planes, "planes_total": len(planes), "planes_enviados": sum(1 for x in planes if x["enviado"]),
            "paradas_total": sum(x["paradas"] for x in planes), "paradas_enviadas": sum(x["paradas_enviadas"] for x in planes)}


def estado_visible(v):
    e = v["estado"]
    if e == "cancelado":
        return "Cancelado"
    if e == "despachado":
        return "Despachado"
    if not v.get("route_id_sr"):
        return "Por enviar a SR" + (" · con camión" if v.get("placa") else "")
    if v.get("placa"):
        return "Con camión"
    return "En SR sin camión"


def pendientes_de_envio(detalle):
    """(viajes sin ruta en SR, paradas pendientes) de un detalle de plan."""
    viajes = paradas = 0
    for v in detalle["viajes"]:
        if v["estado"] in ("cancelado", "despachado"):
            continue
        if not v.get("route_id_sr"):
            viajes += 1
        paradas += sum(1 for p in v["paradas"] if p["estado"] in ("pendiente", "error", "enviando"))
    return viajes, paradas


def sincronizar_despachados(get_conn_fn):
    """Marca como despachado todo viaje cuya ruta de SR ya se importó en Despacho (viaje de Control de Ruta no
    anulado). Es solo SQL, no llama a SR. Devuelve cuántos cambió; nunca lanza."""
    try:
        with closing(get_conn_fn()) as conn:
            sr_int._asegurar_tabla_vinculo(conn)
            with conn.cursor() as cur:
                cur.execute("""
                    UPDATE plan_viaje pv SET estado = 'despachado', id_viaje_cr = s.id_viaje,
                        placa = COALESCE(pv.placa, NULLIF(vj.placa, '')),
                        piloto = COALESCE(pv.piloto, NULLIF(vj.piloto, '')),
                        vehiculo_sr_id = COALESCE(pv.vehiculo_sr_id, (SELECT c.id_sr FROM cat_camiones c WHERE c.placa = vj.placa LIMIT 1)),
                        piloto_sr_id = COALESCE(pv.piloto_sr_id, (SELECT pi.id_sr FROM cat_pilotos pi WHERE pi.nombre = vj.piloto LIMIT 1))
                    FROM sr_vinculo_viaje s JOIN viajes vj ON vj.id_viaje = s.id_viaje
                    WHERE pv.route_id_sr = s.route_id_sr AND vj.estado <> 'Anulado'
                      AND pv.estado IN ('pendiente','en_sr','asignado')""")
                n = cur.rowcount
                # Los ya despachados antes de este cambio y que quedaron sin placa/piloto también se completan.
                cur.execute("""
                    UPDATE plan_viaje pv SET placa = COALESCE(pv.placa, NULLIF(vj.placa, '')), piloto = COALESCE(pv.piloto, NULLIF(vj.piloto, ''))
                    FROM viajes vj WHERE pv.estado = 'despachado' AND pv.id_viaje_cr = vj.id_viaje AND (pv.placa IS NULL OR pv.piloto IS NULL)""")
            conn.commit()
        return n
    except Exception:
        return 0


# ---------------------------------------------------------------------------
# Envío a SimpliRoute
# ---------------------------------------------------------------------------
def _err_sr(data):
    return _corto(data, 400)


def _upd(get_conn_fn, tabla, fila_id, **campos):
    cols = ", ".join(f"{k} = %s" for k in campos)
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute(f"UPDATE {tabla} SET {cols} WHERE id = %s", list(campos.values()) + [fila_id])
        conn.commit()


def _payload_ruta(plan, viaje, cfg, ref, dummy_id):
    dep_nom = viaje.get("deposito_nombre") or cfg.get("deposito_nombre") or "Base Ransa Guatemala"
    lat = viaje.get("deposito_lat") or cfg.get("deposito_lat") or sr_int.UBICACION_BASE_VEHICULOS_SR["latitud"]
    lon = viaje.get("deposito_lon") or cfg.get("deposito_lon") or sr_int.UBICACION_BASE_VEHICULOS_SR["longitud"]
    p = {
        "plan": plan["plan_id_sr"],
        "vehicle": viaje.get("vehiculo_sr_id") or dummy_id,
        "planned_date": plan["fecha"].isoformat(),
        "estimated_time_start": HORA_INICIO_RUTA, "estimated_time_end": HORA_FIN_RUTA,
        "location_start_address": dep_nom, "location_start_latitude": round(float(lat), 6),
        "location_start_longitude": round(float(lon), 6),
        "location_end_address": dep_nom, "location_end_latitude": round(float(lat), 6),
        "location_end_longitude": round(float(lon), 6),
        "reference": ref,
    }
    if viaje.get("piloto_sr_id"):
        p["driver"] = viaje["piloto_sr_id"]          # sin piloto => no se manda el campo (ruta "sin asignar")
    return p


def _titulo_visita(p):
    cod = (p.get("cod_destino") or "").strip()
    return f"{cod} {p['destino']}".strip() if cod.isdigit() else p["destino"]


def visit_type_por_defecto(cliente_orden):
    """GT-DOLLAR CITY -> 'dollarcity' (así viene en SR). Sin él, Despacho no reconocería de qué cliente es la ruta importada."""
    return re.sub(r"[^a-z0-9]", "", _prefijo_codigo(cliente_orden or "").lower())


def _asegurar_visit_type(get_conn_fn, cfg):
    """Si el cliente no tiene visit_type guardado, se guarda el derivado y se deja mapeado al cliente (para la importación de rutas)."""
    vt = cfg.get("visit_type") or visit_type_por_defecto(cfg.get("cliente_orden"))
    if not vt:
        return
    try:
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("UPDATE cat_transporte_cliente SET visit_type = %s WHERE cliente_orden = %s AND (visit_type IS NULL OR visit_type = '')",
                        (vt, cfg["cliente_orden"]))
            cur.execute("INSERT INTO cat_mapeo_cliente_sr (visit_type, cliente) VALUES (%s,%s) ON CONFLICT (visit_type) DO NOTHING", (vt, cfg["cliente"]))
            conn.commit()
    except Exception:
        pass


def _payload_visita(plan, viaje, p, cfg):
    d = {
        "route": viaje["route_id_sr"], "title": _titulo_visita(p),
        "address": p.get("direccion") or p["destino"],
        "latitude": round(float(p["latitud"]), 6), "longitude": round(float(p["longitud"]), 6),
        "planned_date": plan["fecha"].isoformat(), "reference": p["codigo_pedido"], "order": p["orden"],
    }
    if p.get("observaciones"):
        d["notes"] = p["observaciones"]
    vt = cfg.get("visit_type") or visit_type_por_defecto(cfg.get("cliente_orden"))
    if vt:
        d["visit_type"] = vt
    if p.get("duracion"):
        d["duration"] = p["duracion"]
    if p.get("ventana_ini"):
        d["window_start"] = p["ventana_ini"]
    if p.get("ventana_fin"):
        d["window_end"] = p["ventana_fin"]
    return d


def _sin_respuesta(msg):
    """True si el mensaje de error es de red/timeout (SR no llegó a contestar), no un rechazo de SR."""
    m = str(msg)
    return "respondió HTTP" not in m and "limitando las peticiones" not in m


def _crear_visita_sr(payload, token):
    """POST de una visita. Si SR rechaza con 400 un campo que no habíamos probado (duración/ventana, luego
    visit_type) se reintenta sin él y se avisa. Devuelve (ok, visit_id_o_mensaje, nota)."""
    notas = []
    intentos = [payload]
    extras = {k for k in ("duration", "window_start", "window_end") if k in payload}
    if extras:
        intentos.append({k: v for k, v in payload.items() if k not in extras})
    if "visit_type" in payload:
        intentos.append({k: v for k, v in payload.items() if k not in extras and k != "visit_type"})
    ok, data = False, ""
    for i, pl in enumerate(intentos):
        ok, data = sr_int._sr_request("POST", "routes/visits/", json=pl, token=token)
        if ok or "HTTP 400" not in str(data):
            break
        if i + 1 < len(intentos):
            notas.append("SR rechazó un campo opcional y se reintentó sin él")
    if not ok:
        # Sin respuesta de SR (timeout, red caída): la visita pudo haberse creado igual. NO se da por fallida.
        return False, _err_sr(data), ("INCIERTA" if _sin_respuesta(data) else "")
    v = data[0] if isinstance(data, list) else data
    vid = v.get("id") if isinstance(v, dict) else None
    if not vid:
        return False, f"SR no devolvió el id de la visita: {_err_sr(data)}", ""
    nota = ""
    if notas:
        quitados = [k for k in ("duration", "window_start", "window_end", "visit_type")
                    if k in payload and k not in intentos[len(notas)]]
        nota = "Enviada sin " + ", ".join(quitados) + " (SR no los aceptó)."
    return True, vid, nota


def _rutas_sr_del_dia(plan, token):
    """{(plan_id_sr, reference): route_id} de las rutas de SR de la fecha del plan. None si SR no contesta."""
    ok, data = sr_int._sr_request("GET", "routes/routes/", token=token, params={"planned_date": plan["fecha"].isoformat()},
                                  timeout=sr_int.TIMEOUT_LISTA_DRIVERS_SEGUNDOS)
    if not ok:
        return None
    lista = data.get("results", data) if isinstance(data, dict) else data
    return {(str(r.get("plan")), str(r.get("reference"))): str(r.get("id")) for r in (lista or [])}


def _completar_coordenadas(get_conn_fn, plan, viajes, token, res):
    """Antes de crear las visitas: a las paradas sin coordenadas se las busca (1) en el catálogo de tiendas y, si no están, (2) en las visitas
    recientes de SimpliRoute (una sola consulta, que además deja las coordenadas en el catálogo para la próxima vez). Lo que no aparezca queda
    como «error: sin coordenadas» SOLO esa parada (no frena el resto del plan) y se reintenta al volver a enviar."""
    faltan = [(v, p) for v in viajes if v["estado"] not in ("cancelado", "despachado") for p in v["paradas"]
              if p["estado"] in ("pendiente", "error") and (p["latitud"] is None or p["longitud"] is None)]
    if not faltan:
        return

    def _leer_catalogo():
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            for _v, p in faltan:
                if p["latitud"] is not None and p["longitud"] is not None:
                    continue
                cod = str(p.get("cod_destino") or "")
                if cod.isdigit():
                    cur.execute("SELECT latitud, longitud FROM cat_clientes_tiendas WHERE cliente = %s AND codigo_tienda = %s", (plan["cliente"], int(cod)))
                else:
                    cur.execute("SELECT latitud, longitud FROM cat_clientes_tiendas WHERE cliente = %s AND tienda = %s", (plan["cliente"], p["destino"]))
                r = cur.fetchone()
                if r and r[0] is not None and r[1] is not None:
                    cur.execute("UPDATE plan_parada SET latitud = %s, longitud = %s WHERE id = %s", (r[0], r[1], p["id"]))
                    p["latitud"], p["longitud"] = r[0], r[1]
            conn.commit()

    _leer_catalogo()
    if any((p["latitud"] is None or p["longitud"] is None) and p["estado"] == "pendiente" for _v, p in faltan):
        ok, resumen, _det = sr_int.sincronizar_coordenadas_tiendas_desde_sr(get_conn_fn, dias=DIAS_COORD_SR, token=token)
        res["notas"].append(("Coordenadas completadas desde SimpliRoute: " if ok else "No pude consultar las coordenadas en SimpliRoute: ") + str(resumen))
        _leer_catalogo()
    msg = (f"Sin coordenadas: la tienda no aparece en las visitas de SimpliRoute de los últimos {DIAS_COORD_SR} días ni tiene coordenadas en el "
           "catálogo. Cárgalas en Catálogos (columnas latitud y longitud) y vuelve a enviar.")
    for _v, p in faltan:
        if p["latitud"] is None or p["longitud"] is None:
            if p["estado"] != "error" or p.get("error") != msg:
                _upd(get_conn_fn, "plan_parada", p["id"], estado="error", error=msg)
                p["estado"], p["error"] = "error", msg
            res["errores"].append(f"{p['codigo_pedido']} ({p['destino']}): sin coordenadas.")


def _resolver_inciertas(get_conn_fn, plan, viajes, token):
    """Paradas que quedaron 'enviando' (se cortó la conexión a medias): se consulta SR UNA vez por fecha para
    ver si la visita sí se creó, y solo si no existe vuelve a 'pendiente'. Así nunca se duplica una visita."""
    inciertas = [(v, p) for v in viajes for p in v["paradas"] if p["estado"] == "enviando" and v.get("route_id_sr")]
    if not inciertas:
        return 0, ""
    ok, data = sr_int._sr_request("GET", "routes/visits/", token=token, params={"planned_date": plan["fecha"].isoformat()},
                                  timeout=sr_int.TIMEOUT_LISTA_DRIVERS_SEGUNDOS)
    if not ok:
        return 0, f"No pude revisar en SR las paradas que quedaron a medias: {_err_sr(data)}"
    lista = data.get("results", data) if isinstance(data, dict) else data
    por_ref = {(str(x.get("route")), str(x.get("reference"))): x.get("id") for x in (lista or [])}
    arregladas = 0
    for v, p in inciertas:
        vid = por_ref.get((str(v["route_id_sr"]), p["codigo_pedido"]))
        if vid:
            _upd(get_conn_fn, "plan_parada", p["id"], estado="en_sr", visit_id_sr=vid, error=None)
            p["estado"], p["visit_id_sr"] = "en_sr", vid
        else:
            _upd(get_conn_fn, "plan_parada", p["id"], estado="pendiente")
            p["estado"] = "pendiente"
        arregladas += 1
    return arregladas, ""


def enviar_plan_a_sr(get_conn_fn, plan_id, usuario, token=None, max_unidades=None):
    """Envía a SR lo que falte de un plan (plan -> rutas -> visitas), hasta `max_unidades` rutas+visitas por llamada.
    Se puede repetir cuantas veces haga falta: solo manda lo pendiente y guarda cada id de SR apenas responde.
    Devuelve {ok, rutas, paradas, errores:[...], notas:[...], pendientes, mensaje}."""
    max_unidades = max_unidades or LOTE_ENVIO
    res = {"ok": True, "rutas": 0, "paradas": 0, "errores": [], "notas": [], "pendientes": 0, "mensaje": ""}
    if not ENVIO_A_SR_ACTIVO:
        return {**res, "ok": False, "mensaje": "El envío a SR está apagado (ENVIO_A_SR_ACTIVO)."}
    if not token:
        return {**res, "ok": False, "mensaje": "No hay token de SimpliRoute configurado."}
    try:
        d = detalle_plan(get_conn_fn, plan_id)
        if not d:
            return {**res, "ok": False, "mensaje": "El plan ya no existe."}
        plan, viajes = d["plan"], d["viajes"]
        cfgs = [c for c in listar_config(get_conn_fn) if c["cliente_orden"] == plan["cliente_orden"]]
        cfg = cfgs[0] if cfgs else None
        if cfg:
            _asegurar_visit_type(get_conn_fn, cfg)
        if cfg and not pool_propio(get_conn_fn, cfg):
            autoconfigurar_dummy(get_conn_fn, plan["cliente_orden"])
            cfg = next((x for x in listar_config(get_conn_fn) if x["cliente_orden"] == plan["cliente_orden"]), cfg)
        if not cfg:
            return {**res, "ok": False, "mensaje": "Este cliente no está configurado para Transporte (pestaña Configuración)."}
        pool = pool_dummies(get_conn_fn, cfg)          # respaldo: solo para viajes que no traen vehículo en su Placa
        if plan["prueba"] and MARCA_PRUEBA not in plan["nombre"].upper():
            return {**res, "ok": False, "mensaje": "Un plan de prueba debe llevar PRUEBA en el nombre."}

        # 1) El plan en SR (se crea UNA vez; SR no permite borrarlo, por eso se guarda su id).
        if not plan.get("plan_id_sr"):
            ok, data = sr_int._sr_request("POST", "plans/create-plan/", token=token, json={
                "name": plan["nombre"], "start_date": plan["fecha"].isoformat(), "end_date": plan["fecha"].isoformat()})
            if not ok:
                return {**res, "ok": False, "mensaje": "No se pudo crear el plan en SR: " + _err_sr(data) +
                        " Si ya existe un plan con ese nombre en SR, pega su ID en 'Vincular plan existente'."}
            pid = (data[0] if isinstance(data, list) else data or {}).get("id")
            if not pid:
                return {**res, "ok": False, "mensaje": "SR creó el plan pero no devolvió su id: " + _err_sr(data)}
            _upd(get_conn_fn, "plan_transporte", plan_id, plan_id_sr=str(pid))
            plan["plan_id_sr"] = str(pid)
            _registrar(get_conn_fn, usuario, "crear_plan_sr", plan_id, None, plan["nombre"])

        # 2) Paradas que quedaron a medias en un envío anterior.
        _, aviso = _resolver_inciertas(get_conn_fn, plan, viajes, token)
        if aviso:
            res["notas"].append(aviso)

        # 2b) Coordenadas que falten (catálogo → SimpliRoute).
        _completar_coordenadas(get_conn_fn, plan, viajes, token, res)

        # 3) Rutas y visitas pendientes, por lotes.
        unidades = 0
        rutas_dia = None
        for v in viajes:
            if v["estado"] in ("cancelado", "despachado"):
                continue
            pend = [p for p in v["paradas"] if p["estado"] in ("pendiente", "error")]
            if v.get("route_id_sr") and not pend:
                continue
            if unidades >= max_unidades:
                break
            if not v.get("route_id_sr"):
                activas = [p for p in v["paradas"] if p["estado"] != "cancelada"]
                if not activas:
                    continue
                if not any(p["latitud"] is not None and p["longitud"] is not None for p in activas):
                    continue                                    # ninguna parada de este viaje tiene ubicación: no se crea una ruta vacía
                ref = activas[0]["codigo_pedido"]
                rid = None
                dummy_id = v.get("dummy_sr_id") or (pool[(int(v["num"]) - 1) % len(pool)]["id"] if pool else None)
                if str(v.get("error") or "").startswith("INCIERTA"):
                    # La vez anterior SR no contestó: puede que la ruta sí se haya creado. Se revisa antes de crear otra.
                    if rutas_dia is None:
                        rutas_dia = _rutas_sr_del_dia(plan, token)
                    rid = (rutas_dia or {}).get((str(plan["plan_id_sr"]), ref))
                    if rid:
                        res["notas"].append(f"Viaje {v['num']}: la ruta ya existía en SR; se reutilizó.")
                if not rid:
                    dummy_id = v.get("dummy_sr_id") or (pool[(int(v["num"]) - 1) % len(pool)]["id"] if pool else None)
                    if not v.get("vehiculo_sr_id") and not dummy_id:
                        _upd(get_conn_fn, "plan_viaje", v["id"], error="Falta el vehículo (columna Placa del Excel)")
                        res["errores"].append(f"Viaje {v['num']}: falta el vehículo (columna Placa del Excel).")
                        continue
                    payload = _payload_ruta(plan, v, cfg, ref, dummy_id)
                    ok, data = sr_int._sr_request("POST", "routes/routes/", token=token, json=payload)
                    unidades += 1
                    if not ok:
                        marca = "INCIERTA: " if _sin_respuesta(data) else ""
                        _upd(get_conn_fn, "plan_viaje", v["id"], error=marca + _err_sr(data))
                        res["errores"].append(f"Viaje {v['num']}: no se pudo crear la ruta. {_err_sr(data)}")
                        continue
                    r = data[0] if isinstance(data, list) else data
                    rid = (r or {}).get("id")
                if not rid:
                    _upd(get_conn_fn, "plan_viaje", v["id"], error="SR no devolvió el id de la ruta")
                    res["errores"].append(f"Viaje {v['num']}: SR no devolvió el id de la ruta.")
                    continue
                nuevo_estado = "asignado" if v.get("vehiculo_sr_id") else "en_sr"
                _upd(get_conn_fn, "plan_viaje", v["id"], route_id_sr=str(rid), estado=nuevo_estado, error=None,
                     dummy_sr_id=(None if v.get("vehiculo_sr_id") else dummy_id))
                v["route_id_sr"], v["estado"] = str(rid), nuevo_estado
                res["rutas"] += 1
                pend = [p for p in v["paradas"] if p["estado"] in ("pendiente", "error")]
            for p in pend:
                if unidades >= max_unidades:
                    break
                if p["latitud"] is None or p["longitud"] is None:
                    continue                                    # ya quedó marcada «sin coordenadas» arriba (_completar_coordenadas)
                _upd(get_conn_fn, "plan_parada", p["id"], estado="enviando")        # antes de llamar a SR
                ok, valor, nota = _crear_visita_sr(_payload_visita(plan, v, p, cfg), token)
                unidades += 1
                if ok:
                    _upd(get_conn_fn, "plan_parada", p["id"], estado="en_sr", visit_id_sr=int(valor), error=None)
                    res["paradas"] += 1
                    if nota:
                        res["notas"].append(f"{p['codigo_pedido']}: {nota}")
                elif nota == "INCIERTA":
                    _upd(get_conn_fn, "plan_parada", p["id"], error=valor)    # sigue 'enviando': se verifica en SR al reintentar
                    res["errores"].append(f"{p['codigo_pedido']}: SR no respondió; quedó por verificar (al reenviar se revisa en SR "
                                          f"si ya se había creado, para no duplicarla).")
                else:
                    _upd(get_conn_fn, "plan_parada", p["id"], estado="error", error=valor)
                    res["errores"].append(f"{p['codigo_pedido']}: {valor}")
        d2 = detalle_plan(get_conn_fn, plan_id)
        vp, pp = pendientes_de_envio(d2)
        res["pendientes"] = vp + pp
        _registrar(get_conn_fn, usuario, "enviar_sr", plan_id, None,
                   f"{res['rutas']} ruta(s), {res['paradas']} visita(s), {len(res['errores'])} error(es)")
        if res["errores"]:
            res["ok"] = False
        res["mensaje"] = (f"Enviado a SR: {res['rutas']} ruta(s) y {res['paradas']} visita(s)."
                          + (f" Quedan {res['pendientes']} por enviar." if res["pendientes"] else ""))
        return res
    except Exception as e:
        return {**res, "ok": False, "mensaje": f"Error inesperado al enviar a SR ({type(e).__name__}).", "excepcion": e}


_PALABRAS_NO_CLIENTE = {"CITY", "LOCAL", "GUATEMALA", "SOCIEDAD", "ANONIMA", "COMERCIAL", "DE", "LA", "LOS", "LAS", "SECO", "FRIO"}


def _palabras_cliente(cliente_orden, cliente):
    """Palabras que identifican al cliente en el nombre de su dummy: 'GT-DOLLAR CITY' -> {'DOLLAR'}; 'Diana (La Comercial)' -> {'DIANA'}."""
    palabras = set()
    for t in (_prefijo_codigo(cliente_orden), re.sub(r"\(.*?\)", " ", cliente or "")):
        for w in re.findall(r"[A-Z0-9]+", _norm(t).upper()):
            if len(w) >= 4 and w not in _PALABRAS_NO_CLIENTE:
                palabras.add(w)
    return palabras


def vehiculos_sincronizados(get_conn_fn, solo_dummies=True):
    """Vehículos de cat_camiones con id de SR. solo_dummies=True deja fuera las placas normales (123ABC): lo que queda son
    los provisionales con nombre (DOLLARCITY1, UNISUPERIMPORTADO25...), que la sincronización de camiones ya trajo de SR."""
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT placa, id_sr FROM cat_camiones WHERE id_sr IS NOT NULL ORDER BY placa")
        filas = cur.fetchall()
    out = [{"id": int(i), "nombre": p} for p, i in filas]
    if solo_dummies:
        out = [v for v in out if not re.match(r"^[0-9]{3}[A-Z]{3}$", str(v["nombre"]).upper())]
    return out


def es_generico(nombre):
    """Vehículo genérico (DC01, ADICIONALES1 SECO, UNISUPERIMPORTADO25...): nombre con letras y número, no una placa normal (896BYM)."""
    k = sr_int.clave_placa(nombre)
    return bool(k) and not re.match(r"^[0-9]{3}[A-Z]{3}$", k) and bool(re.match(r"^[A-Z]{2,}[0-9]+[A-Z]*$", k))


def _clave_dummy(nombre):
    return re.sub(r"[^A-Z0-9]", "", str(nombre).upper())


def grupos_dummy(get_conn_fn):
    """Grupos de dummies ya sincronizados, por prefijo: {'DC': [{'id','nombre','num'}...], 'DI': [...]}. Un dummy es un vehículo con
    nombre letras+número (DC01, DC02...); las placas normales (896BYM) no cuentan."""
    grupos = {}
    for v in vehiculos_sincronizados(get_conn_fn):
        m = re.match(r"^([A-Z]+)(\d+)$", _clave_dummy(v["nombre"]))
        if m and not m.group(1).startswith(DUMMY_GENERICO):
            grupos.setdefault(m.group(1), []).append(dict(v, num=int(m.group(2))))
    for lista in grupos.values():
        lista.sort(key=lambda x: x["num"])
    return grupos


def dummies_genericos(get_conn_fn):
    """Vehículos genéricos «ADICIONALES…» ya sincronizados, en orden (ADICIONALES1 SECO, ADICIONALES2 SECO…). Lista de {'id','nombre'}."""
    out = [v for v in vehiculos_sincronizados(get_conn_fn) if _clave_dummy(v["nombre"]).startswith(DUMMY_GENERICO)]
    def _clave_orden(v):
        m = re.search(r"(\d+)", str(v["nombre"]))
        return (int(m.group(1)) if m else 0, "FRIO" in _clave_dummy(v["nombre"]), str(v["nombre"]))   # primero los SECO
    return sorted(out, key=_clave_orden)


def pool_propio(get_conn_fn, cfg):
    """Dummies PROPIOS del cliente: su grupo configurado (DC01, DC02...) o, si no hay grupo, el único vehículo elegido."""
    pref = _clave_dummy(cfg.get("dummy_prefijo") or "")
    if pref:
        g = grupos_dummy(get_conn_fn).get(pref, [])
        if g:
            return [{"id": x["id"], "nombre": x["nombre"]} for x in g]
    if cfg.get("vehiculo_dummy_sr"):
        return [{"id": int(cfg["vehiculo_dummy_sr"]), "nombre": cfg.get("vehiculo_dummy_nombre") or ""}]
    return []


def pool_dummies(get_conn_fn, cfg):
    """Dummies que usan los viajes de este cliente: los propios si tiene; si no, los genéricos ADICIONALES."""
    return pool_propio(get_conn_fn, cfg) or [{"id": v["id"], "nombre": v["nombre"]} for v in dummies_genericos(get_conn_fn)]


def todos_los_dummies_ids(get_conn_fn, cfg):
    """Ids de cualquier vehículo que cuente como «sin camión» para este cliente (propios o genéricos)."""
    return {x["id"] for x in pool_propio(get_conn_fn, cfg)} | {v["id"] for v in dummies_genericos(get_conn_fn)}


def sugerir_prefijo(cliente_orden, grupos):
    """DOLLAR CITY -> 'DC' (iniciales) si existe ese grupo. Solo para nombres de 2 palabras o más; si no, no adivina."""
    palabras = [w for w in re.findall(r"[A-Z0-9]+", _norm(_prefijo_codigo(cliente_orden)).upper()) if w not in _PALABRAS_NO_CLIENTE or w == "CITY"]
    if len(palabras) >= 2:
        ini = "".join(w[0] for w in palabras)
        if ini in grupos:
            return ini
    return None


def candidatos_dummy(get_conn_fn, cliente_orden, cliente):
    """Para clientes con UN solo dummy cuyo nombre contiene el del cliente (ej. UNISUPERIMPORTADO25)."""
    palabras = _palabras_cliente(cliente_orden, cliente)
    if not palabras:
        return []
    return [v for v in vehiculos_sincronizados(get_conn_fn)
            if any(w in _clave_dummy(v["nombre"]) for w in palabras) and not _clave_dummy(v["nombre"]).startswith(DUMMY_GENERICO)]


def autoconfigurar_dummy(get_conn_fn, cliente_orden, usuario="sistema"):
    """Deja configurado el dummy del cliente si se puede identificar sin ambigüedad: (1) por las iniciales (DOLLAR CITY -> grupo DC) o
    (2) un único vehículo con su nombre. Devuelve (ok, texto|n_candidatos)."""
    try:
        cfgs = [c for c in listar_config(get_conn_fn) if c["cliente_orden"] == cliente_orden]
        if not cfgs:
            return False, 0
        c = cfgs[0]
        pool = pool_propio(get_conn_fn, c)
        if pool:
            return True, (c.get("dummy_prefijo") or pool[0]["nombre"])
        grupos = grupos_dummy(get_conn_fn)
        pref = sugerir_prefijo(cliente_orden, grupos)
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            if pref:
                cur.execute("UPDATE cat_transporte_cliente SET dummy_prefijo=%s WHERE cliente_orden=%s AND dummy_prefijo IS NULL", (pref, cliente_orden))
                _bitacora(cur, usuario, "dummy_automatico", detalle=f"{cliente_orden} -> grupo {pref}")
                conn.commit()
                return True, pref
        cands = candidatos_dummy(get_conn_fn, cliente_orden, c["cliente"])
        if len(cands) != 1:
            return False, len(cands)
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("UPDATE cat_transporte_cliente SET vehiculo_dummy_sr=%s, vehiculo_dummy_nombre=%s WHERE cliente_orden=%s AND vehiculo_dummy_sr IS NULL",
                        (cands[0]["id"], cands[0]["nombre"], cliente_orden))
            _bitacora(cur, usuario, "dummy_automatico", detalle=f"{cliente_orden} -> {cands[0]['nombre']}")
            conn.commit()
        return True, cands[0]["nombre"]
    except Exception:
        return False, 0


def _cliente_cr_por_nombre(cur, cliente_orden):
    """Cliente de Control de Ruta que corresponde a un «Cliente Orden» de la plantilla, si hay UNO solo posible por nombre
    (GT-DOLLAR CITY -> el único cliente cuyo nombre contiene DOLLAR). None si no hay o hay varios."""
    palabras = _palabras_cliente(cliente_orden, "")
    if not palabras:
        return None
    cur.execute("SELECT nombre FROM cat_clientes WHERE COALESCE(activo, TRUE)")
    hallados = [r[0] for r in cur.fetchall() if any(w in re.sub(r"[^A-Z0-9]", "", _norm(r[0]).upper()) for w in palabras)]
    return hallados[0] if len(hallados) == 1 else None


def requisitos_envio(get_conn_fn, planes, token_fn):
    """Qué falta para poder enviar estos planes a SR. Devuelve lista de textos (vacía = todo listo)."""
    problemas = []
    try:
        tok = token_fn()
    except Exception:
        tok = None
    if not tok:
        problemas.append("No hay token de SimpliRoute disponible (se toma de Supabase Vault, secreto «simpliroute_api_token»). "
                         "Avísale a un Administrador.")
    cfg = {c["cliente_orden"]: c for c in listar_config(get_conn_fn)}
    for co in sorted({pl["cliente_orden"] for pl in planes}):
        if co not in cfg:
            problemas.append(f"{co}: no está configurado para Transporte (pestaña Configuración).")
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        for pl in planes:
            c = cfg.get(pl["cliente_orden"])
            if c and pool_dummies(get_conn_fn, c):
                continue                                     # tiene dummy por defecto configurado: cubre a los viajes sin Placa
            cur.execute("""SELECT COUNT(*) FROM plan_viaje WHERE plan_id = %s AND route_id_sr IS NULL AND estado NOT IN ('cancelado','despachado')
                           AND vehiculo_sr_id IS NULL AND dummy_sr_id IS NULL""", (pl["id"],))
            n = cur.fetchone()[0]
            if n:
                problemas.append(f"{pl['nombre']}: {n} viaje(s) sin vehículo. Falta la columna Placa en el Excel (usa el genérico, "
                                 "como DC01 o ADICIONALES1, o el camión real).")
    return problemas


def advertencias_envio(get_conn_fn, planes):
    """Avisos que NO bloquean el envío pero conviene leer: p. ej. depósito sin coordenadas (las rutas saldrían de la ubicación base)."""
    avisos = []
    try:
        cfg = {c["cliente_orden"]: c for c in listar_config(get_conn_fn)}
        for co in sorted({pl["cliente_orden"] for pl in planes}):
            c = cfg.get(co)
            if c and (c.get("deposito_lat") is None or c.get("deposito_lon") is None):
                avisos.append(f"{co}: el depósito no tiene coordenadas en Configuración; las rutas saldrán y regresarán a la ubicación base de Ransa, "
                              "no al depósito, y SR calculará mal sus distancias.")
    except Exception:
        pass
    return avisos


PRESUPUESTO_ENVIO_SEGUNDOS = 100     # tiempo máximo de un clic de "enviar todos"; si falta algo, se avisa y se vuelve a presionar


def enviar_planes(get_conn_fn, plan_ids, usuario, token=None, presupuesto=None, max_unidades=None):
    """Envía a SR varios planes seguidos (todos los de una fecha). Por cada plan repite los lotes hasta terminar, se
    detiene si no hay avance o si se acaba el tiempo, y un plan con errores NO frena a los demás.
    Devuelve (resumen, ok): resumen = lista de dicts {plan_id, plan, cliente, cliente_orden, plan_nuevo_en_sr, rutas,
    paradas, pendientes, errores, notas, mensaje}."""
    presupuesto = PRESUPUESTO_ENVIO_SEGUNDOS if presupuesto is None else presupuesto
    t0 = _time.time()
    resumen, todo_ok = [], True
    for pid in plan_ids:
        d = detalle_plan(get_conn_fn, pid)
        if not d:
            continue
        plan = d["plan"]
        fila = {"plan_id": pid, "plan": plan["nombre"], "cliente": plan["cliente"], "cliente_orden": plan["cliente_orden"],
                "plan_nuevo_en_sr": False, "rutas": 0, "paradas": 0, "pendientes": 0, "errores": [], "notas": [], "mensaje": ""}
        vp, pp = pendientes_de_envio(d)
        fila["pendientes"] = vp + pp
        if not plan.get("plan_id_sr") and (vp + pp) > 0:
            fila["plan_nuevo_en_sr"] = True
        if (vp + pp) == 0:
            fila["mensaje"] = "Ya estaba todo enviado."
            resumen.append(fila)
            continue
        while True:
            if _time.time() - t0 > presupuesto:
                fila["mensaje"] = "Se acabó el tiempo de este envío; presiona de nuevo para continuar."
                break
            r = enviar_plan_a_sr(get_conn_fn, pid, usuario, token=token, max_unidades=max_unidades)
            fila["rutas"] += r["rutas"]
            fila["paradas"] += r["paradas"]
            fila["pendientes"] = r["pendientes"]
            fila["errores"] = list(dict.fromkeys(fila["errores"] + r["errores"]))      # sin repetir el mismo error en cada lote
            fila["notas"] = list(dict.fromkeys(fila["notas"] + r["notas"]))
            fila["mensaje"] = r["mensaje"]
            if not r["ok"] and r["rutas"] == 0 and r["paradas"] == 0:
                break                                   # sin avance: no tiene caso repetir
            if r["pendientes"] == 0:
                break
        if fila["errores"] or fila["pendientes"]:
            todo_ok = False
        if fila["plan_nuevo_en_sr"]:
            with closing(get_conn_fn()) as conn, conn.cursor() as cur:
                cur.execute("SELECT plan_id_sr FROM plan_transporte WHERE id = %s", (pid,))
                x = cur.fetchone()
                fila["plan_nuevo_en_sr"] = bool(x and x[0])
        resumen.append(fila)
    _registrar(get_conn_fn, usuario, "enviar_todos", None, None,
               "; ".join(f"{f['cliente']}: {f['rutas']} ruta(s), {f['paradas']} visita(s), {f['pendientes']} pendiente(s)" for f in resumen))
    return resumen, todo_ok


def _registrar(get_conn_fn, usuario, accion, plan_id=None, viaje_id=None, detalle=""):
    try:
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            _bitacora(cur, usuario, accion, plan_id, viaje_id, detalle)
            conn.commit()
    except Exception:
        pass


def vincular_plan_existente(get_conn_fn, plan_id, plan_id_sr, usuario):
    """Para cuando SR ya tiene un plan con ese nombre (no se puede borrar): se pega su id para reutilizarlo."""
    plan_id_sr = _txt(plan_id_sr)
    if not re.match(r"^[0-9a-fA-F-]{20,}$", plan_id_sr):
        return False, "El ID del plan en SR no parece válido (es un código largo con guiones)."
    try:
        _upd(get_conn_fn, "plan_transporte", plan_id, plan_id_sr=plan_id_sr)
        _registrar(get_conn_fn, usuario, "vincular_plan_sr", plan_id, None, plan_id_sr)
        return True, "Plan vinculado con SR."
    except Exception as e:
        return False, f"No se pudo vincular ({type(e).__name__})."


# ---------------------------------------------------------------------------
# Asignación de camión y piloto
# ---------------------------------------------------------------------------
def _fijar_camion_piloto(cur, viaje_id, placa, piloto, usuario, solo_local=False):
    cur.execute("SELECT id_sr FROM cat_camiones WHERE placa = %s", (placa,))
    r = cur.fetchone()
    veh_id = r[0] if r else None
    pil_id = None
    if piloto:
        cur.execute("SELECT id_sr FROM cat_pilotos WHERE nombre = %s", (piloto,))
        r2 = cur.fetchone()
        pil_id = r2[0] if r2 else None
    cur.execute("UPDATE plan_viaje SET placa=%s, piloto=%s, vehiculo_sr_id=%s, piloto_sr_id=%s, asignado_por=%s, "
                "asignado_en=NOW(), cambiado_en_sr=FALSE WHERE id=%s", (placa, piloto, veh_id, pil_id, usuario, viaje_id))


def avisos_doble_uso(get_conn_fn, viaje_id, placa, piloto, fecha):
    """Avisos (no bloqueos) si la placa o el piloto ya tienen otro viaje ese día: en otros planes o en Despacho."""
    avisos = []
    try:
        clave = sr_int.clave_placa(placa)
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("""
                SELECT p.nombre, v.num FROM plan_viaje v JOIN plan_transporte p ON p.id = v.plan_id
                WHERE v.id <> %s AND p.fecha = %s AND v.estado <> 'cancelado'
                  AND REPLACE(REPLACE(UPPER(COALESCE(v.placa,'')),'-',''),' ','') LIKE %s""",
                        (viaje_id, fecha, "%" + clave))
            for nom, num in cur.fetchall():
                avisos.append(f"La placa {placa} ya está asignada al viaje {num} del plan {nom}.")
            if piloto:
                cur.execute("""
                    SELECT p.nombre, v.num FROM plan_viaje v JOIN plan_transporte p ON p.id = v.plan_id
                    WHERE v.id <> %s AND p.fecha = %s AND v.estado <> 'cancelado' AND v.piloto = %s""",
                            (viaje_id, fecha, piloto))
                for nom, num in cur.fetchall():
                    avisos.append(f"El piloto {piloto} ya está asignado al viaje {num} del plan {nom}.")
            try:
                cur.execute("""
                    SELECT id_viaje, cliente FROM viajes
                    WHERE fecha_creacion = %s AND estado <> 'Anulado'
                      AND REPLACE(REPLACE(UPPER(placa),'-',''),' ','') LIKE %s""", (fecha.isoformat(), "%" + clave))
                for idv, cli in cur.fetchall():
                    avisos.append(f"La placa {placa} ya tiene el viaje {idv} ({cli}) en Despacho ese día.")
            except Exception:
                conn.rollback()
    except Exception:
        pass
    return avisos


def asignar_viaje(get_conn_fn, viaje_id, placa, piloto, usuario, token=None):
    """Asigna (o cambia) camión y piloto de un viaje. Si la ruta ya está en SR, primero se cambia allá (PATCH, con
    confirmación posterior) y SOLO si SR lo acepta se guarda aquí. Devuelve (ok, mensaje, avisos)."""
    try:
        d = _viaje_y_plan(get_conn_fn, viaje_id)
        if not d:
            return False, "El viaje ya no existe.", []
        v, plan = d
        if v["estado"] in ("despachado", "cancelado"):
            return False, f"El viaje está {v['estado']}: ya no se puede cambiar su asignación.", []
        if not placa:
            return False, "Elige una placa.", []
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("SELECT placa, id_sr, sincronizado_sr FROM cat_camiones WHERE placa = %s", (placa,))
            cam = cur.fetchone()
            if not cam:
                return False, "Esa placa no existe en el catálogo de camiones.", []
            if not cam[1]:
                return False, f"El camión {placa} no está sincronizado con SR; sincronízalo en Catálogos antes de asignarlo.", []
            pil_id = None
            if piloto:
                cur.execute("SELECT id_sr FROM cat_pilotos WHERE nombre = %s", (piloto,))
                r = cur.fetchone()
                if not r:
                    return False, "Ese piloto no existe en el catálogo.", []
                pil_id = r[0]
        avisos = avisos_doble_uso(get_conn_fn, viaje_id, placa, piloto, plan["fecha"])
        if piloto and not pil_id:
            avisos.append(f"El piloto {piloto} no está sincronizado con SR: en SR solo cambiará el camión.")
        msg_sr = ""
        if v.get("route_id_sr"):
            if not token:
                return False, "No hay token de SimpliRoute: no se puede cambiar la ruta en SR.", avisos
            ok, msg_sr = sr_int.reasignar_vehiculo_piloto_ruta_sr(v["route_id_sr"], cam[1], pil_id, token=token)
            if not ok:
                return False, f"SR no aceptó el cambio: {_corto(msg_sr, 300)}", avisos
        with closing(get_conn_fn()) as conn:
            try:
                with conn.cursor() as cur:
                    _fijar_camion_piloto(cur, viaje_id, placa, piloto or None, usuario)
                    if v.get("route_id_sr"):
                        cur.execute("UPDATE plan_viaje SET estado='asignado', error=NULL WHERE id=%s", (viaje_id,))
                    _bitacora(cur, usuario, "asignar", plan["id"], viaje_id, f"{placa} · {piloto or 'sin piloto'}")
                conn.commit()
            except Exception:
                conn.rollback()
                raise
        return True, ("Asignado en Control de Ruta y SimpliRoute." if v.get("route_id_sr")
                      else "Asignado. Se aplicará en SR cuando se envíe el plan."), avisos
    except Exception as e:
        return False, f"No se pudo asignar ({type(e).__name__}).", []


def _viaje_y_plan(get_conn_fn, viaje_id):
    import psycopg2.extras
    with closing(get_conn_fn()) as conn, conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
        cur.execute("SELECT * FROM plan_viaje WHERE id = %s", (viaje_id,))
        v = cur.fetchone()
        if not v:
            return None
        cur.execute("SELECT * FROM plan_transporte WHERE id = %s", (v["plan_id"],))
        p = cur.fetchone()
        return dict(v), dict(p)


def refrescar_desde_sr(get_conn_fn, plan_id, usuario, token=None, solo_completar=False):
    """Consulta las rutas de SR de la fecha del plan (sin visitas: es ligero) y detecta lo que se cambió allá:
    otro camión/piloto, o una ruta que ya no existe. Nunca pisa en silencio: marca el viaje como 'cambiado en SR'.
    Con solo_completar=True (lo que se hace solo al abrir la pantalla) únicamente LLENA lo que Control de Ruta no tiene (camión o piloto
    vacíos) con lo que figure en SR; lo que Control de Ruta ya tiene no se toca. Devuelve (ok, mensaje, cambios)."""
    if not token:
        return False, "No hay token de SimpliRoute.", []
    try:
        d = detalle_plan(get_conn_fn, plan_id)
        if not d:
            return False, "El plan ya no existe.", []
        plan, viajes = d["plan"], d["viajes"]
        con_ruta = [v for v in viajes if v.get("route_id_sr") and v["estado"] in ("en_sr", "asignado")]
        if not con_ruta:
            return True, "No hay rutas en SR para revisar todavía.", []
        ok, data = sr_int._sr_request("GET", "routes/routes/", token=token, params={"planned_date": plan["fecha"].isoformat()},
                                      timeout=sr_int.TIMEOUT_LISTA_DRIVERS_SEGUNDOS)
        if not ok:
            return False, f"No pude consultar SR: {_err_sr(data)}", []
        lista = data.get("results", data) if isinstance(data, dict) else data
        rutas = {str(r.get("id")): r for r in (lista or [])}
        cfgs = [c for c in listar_config(get_conn_fn) if c["cliente_orden"] == plan["cliente_orden"]]
        ids_dummy = (todos_los_dummies_ids(get_conn_fn, cfgs[0]) if cfgs else set()) | \
            {x["id"] for x in vehiculos_sincronizados(get_conn_fn) if es_generico(x["nombre"])}
        cambios = []
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("SELECT placa, id_sr FROM cat_camiones WHERE id_sr IS NOT NULL")
            placa_de = {int(i): p for p, i in cur.fetchall()}
            cur.execute("SELECT nombre, id_sr FROM cat_pilotos WHERE id_sr IS NOT NULL")
            piloto_de = {int(i): n for n, i in cur.fetchall()}
            for v in con_ruta:
                r = rutas.get(str(v["route_id_sr"]))
                if r is None:
                    cur.execute("UPDATE plan_viaje SET error=%s WHERE id=%s",
                                ("La ruta no se encontró en SR (¿se borró allá?)", v["id"]))
                    if not solo_completar:
                        cambios.append(f"Viaje {v['num']}: la ruta ya no existe en SR.")
                    continue
                veh_sr, drv_sr = r.get("vehicle"), r.get("driver")
                if solo_completar:
                    sets, vals, partes = [], [], []
                    if v.get("vehiculo_sr_id") is None and veh_sr is not None and veh_sr not in ids_dummy:
                        placa = placa_de.get(int(veh_sr))
                        sets += ["placa = %s", "vehiculo_sr_id = %s", "estado = 'asignado'", "cambiado_en_sr = TRUE"]
                        vals += [placa, veh_sr]
                        partes.append(placa or "un vehículo que no está en el catálogo")
                    if v.get("piloto_sr_id") is None and not v.get("piloto") and drv_sr is not None:
                        piloto = piloto_de.get(int(drv_sr))
                        sets += ["piloto = %s", "piloto_sr_id = %s"]
                        vals += [piloto, drv_sr]
                        partes.append(piloto or "un piloto que no está en el catálogo")
                    if sets:
                        cur.execute(f"UPDATE plan_viaje SET {', '.join(sets)} WHERE id = %s", vals + [v["id"]])
                        cambios.append(f"Viaje {v['num']}: " + " / ".join(partes) + ".")
                    continue
                mismo_veh = (veh_sr == v.get("vehiculo_sr_id")) or (v.get("vehiculo_sr_id") is None and veh_sr in ids_dummy)
                mismo_drv = (drv_sr == v.get("piloto_sr_id")) or (drv_sr is None and v.get("piloto_sr_id") is None)
                if mismo_veh and mismo_drv:
                    continue
                if veh_sr in ids_dummy:
                    cur.execute("UPDATE plan_viaje SET placa=NULL, piloto=NULL, vehiculo_sr_id=NULL, piloto_sr_id=NULL, "
                                "estado='en_sr', cambiado_en_sr=TRUE WHERE id=%s", (v["id"],))
                    cambios.append(f"Viaje {v['num']}: en SR se volvió a dejar sin camión.")
                else:
                    placa = placa_de.get(int(veh_sr)) if veh_sr is not None else None
                    piloto = piloto_de.get(int(drv_sr)) if drv_sr is not None else None
                    cur.execute("UPDATE plan_viaje SET placa=%s, piloto=%s, vehiculo_sr_id=%s, piloto_sr_id=%s, "
                                "estado='asignado', cambiado_en_sr=TRUE WHERE id=%s",
                                (placa, piloto, veh_sr, drv_sr, v["id"]))
                    cambios.append(f"Viaje {v['num']}: cambiado en SR a {placa or 'un vehículo que no está en el catálogo'}"
                                   + (f" / {piloto}" if piloto else "") + ".")
            if cambios:
                _bitacora(cur, usuario, "refrescar_sr", plan_id, None, "; ".join(cambios))
            conn.commit()
        return True, ("Revisado contra SR." + (f" {len(cambios)} cambio(s) detectado(s)." if cambios else " Todo coincide.")), cambios
    except Exception as e:
        return False, f"No se pudo revisar contra SR ({type(e).__name__}).", []


# ---------------------------------------------------------------------------
# Cancelar viajes y paradas
# ---------------------------------------------------------------------------
def _borrar_sr(ruta, token):
    ok, data = sr_int._sr_request("DELETE", ruta, token=token)
    if ok or "HTTP 404" in str(data):
        return True, ""
    return False, _err_sr(data)


def cancelar_parada(get_conn_fn, parada_id, usuario, token=None):
    """Cancela una parada: si ya está en SR se borra su visita allá. Si era la última, el viaje también se cancela."""
    try:
        import psycopg2.extras
        with closing(get_conn_fn()) as conn, conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute("SELECT pa.*, v.estado AS estado_viaje, v.route_id_sr FROM plan_parada pa "
                        "JOIN plan_viaje v ON v.id = pa.viaje_id WHERE pa.id = %s", (parada_id,))
            p = cur.fetchone()
        if not p:
            return False, "La parada ya no existe."
        if p["estado_viaje"] in ("despachado", "cancelado"):
            return False, f"El viaje está {p['estado_viaje']}: no se pueden quitar paradas."
        if p["estado"] == "cancelada":
            return True, "La parada ya estaba cancelada."
        if p["visit_id_sr"]:
            if not token:
                return False, "No hay token de SimpliRoute: no se puede borrar la visita en SR."
            ok, msg = _borrar_sr(f"routes/visits/{p['visit_id_sr']}/", token)
            if not ok:
                return False, f"SR no dejó borrar la visita: {msg}"
        _upd(get_conn_fn, "plan_parada", parada_id, estado="cancelada", error=None)
        _registrar(get_conn_fn, usuario, "cancelar_parada", p["plan_id"], p["viaje_id"], p["codigo_pedido"])
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("SELECT COUNT(*) FROM plan_parada WHERE viaje_id=%s AND estado <> 'cancelada'", (p["viaje_id"],))
            quedan = cur.fetchone()[0]
        if quedan == 0:
            ok, msg = cancelar_viaje(get_conn_fn, p["viaje_id"], usuario, token=token)
            return ok, ("Parada cancelada; era la última, por eso el viaje también se canceló." if ok else msg)
        return True, "Parada cancelada."
    except Exception as e:
        return False, f"No se pudo cancelar ({type(e).__name__})."


def cancelar_viaje(get_conn_fn, viaje_id, usuario, token=None):
    """Cancela un viaje: en SR se borran primero sus visitas y luego la ruta (en ese orden). Si SR falla a la mitad,
    lo ya borrado queda registrado y el viaje NO se marca cancelado, para poder reintentar."""
    try:
        d = _viaje_y_plan(get_conn_fn, viaje_id)
        if not d:
            return False, "El viaje ya no existe."
        v, plan = d
        if v["estado"] == "despachado":
            return False, "El viaje ya fue despachado."
        if v["estado"] == "cancelado":
            return True, "El viaje ya estaba cancelado."
        import psycopg2.extras
        with closing(get_conn_fn()) as conn, conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute("SELECT id, visit_id_sr, estado FROM plan_parada WHERE viaje_id = %s", (viaje_id,))
            paradas = [dict(r) for r in cur.fetchall()]
        necesita_sr = bool(v.get("route_id_sr")) or any(p["visit_id_sr"] and p["estado"] != "cancelada" for p in paradas)
        if necesita_sr and not token:
            return False, "No hay token de SimpliRoute: no se puede borrar la ruta en SR."
        for p in paradas:
            if p["visit_id_sr"] and p["estado"] != "cancelada":
                ok, msg = _borrar_sr(f"routes/visits/{p['visit_id_sr']}/", token)
                if not ok:
                    return False, f"SR no dejó borrar una visita ({msg}). El viaje sigue activo; puedes reintentar."
                _upd(get_conn_fn, "plan_parada", p["id"], estado="cancelada")
        if v.get("route_id_sr"):
            ok, msg = _borrar_sr(f"routes/routes/{v['route_id_sr']}/", token)
            if not ok:
                return False, f"SR no dejó borrar la ruta ({msg}). El viaje sigue activo; puedes reintentar."
        with closing(get_conn_fn()) as conn:
            try:
                with conn.cursor() as cur:
                    cur.execute("UPDATE plan_parada SET estado='cancelada' WHERE viaje_id=%s", (viaje_id,))
                    cur.execute("UPDATE plan_viaje SET estado='cancelado', error=NULL WHERE id=%s", (viaje_id,))
                    _bitacora(cur, usuario, "cancelar_viaje", plan["id"], viaje_id, f"Viaje {v['num']}")
                conn.commit()
            except Exception:
                conn.rollback()
                raise
        return True, "Viaje cancelado."
    except Exception as e:
        return False, f"No se pudo cancelar ({type(e).__name__})."


# ---------------------------------------------------------------------------
# Seguimiento (reporte): solicitados vs planeados vs con camión
# ---------------------------------------------------------------------------
def resumen_seguimiento(get_conn_fn, fecha_ini, fecha_fin, clientes=None, prueba=False):
    """Una fila por (fecha, cliente de Control de Ruta). `clientes` = lista de nombres permitidos (None = todos).
    Faltan = solicitados - (con camión + despachados), nunca menor que 0. Sin solicitud registrada => None."""
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        filtro = ""
        params = [bool(prueba), fecha_ini, fecha_fin]
        if clientes is not None:
            filtro = " AND p.cliente = ANY(%s)"
            params.append(list(clientes))
        cur.execute(f"""
            SELECT p.fecha, p.cliente,
              COUNT(v.id) FILTER (WHERE v.estado <> 'cancelado') AS planeados,
              COUNT(v.id) FILTER (WHERE v.estado = 'pendiente') AS por_enviar,
              COUNT(v.id) FILTER (WHERE v.estado = 'en_sr') AS sin_camion,
              COUNT(v.id) FILTER (WHERE v.estado = 'asignado') AS con_camion,
              COUNT(v.id) FILTER (WHERE v.estado = 'despachado') AS despachados
            FROM plan_transporte p LEFT JOIN plan_viaje v ON v.plan_id = p.id
            WHERE p.prueba = %s AND p.fecha BETWEEN %s AND %s {filtro}
            GROUP BY p.fecha, p.cliente""", params)
        filas = {(f, c): dict(fecha=f, cliente=c, planeados=pl, por_enviar=pe, sin_camion=sc, con_camion=cc,
                              despachados=de, solicitados=None)
                 for f, c, pl, pe, sc, cc, de in cur.fetchall()}
        cur.execute("SELECT s.fecha, s.cliente, s.vehiculos_solicitados FROM solicitudes_transporte s "
                    "WHERE s.fecha BETWEEN %s AND %s AND s.cliente IN (SELECT DISTINCT cliente FROM cat_transporte_cliente WHERE activo)"
                    + (" AND s.cliente = ANY(%s)" if clientes is not None else ""),
                    [fecha_ini, fecha_fin] + ([list(clientes)] if clientes is not None else []))
        for f, c, n in cur.fetchall():
            fila = filas.setdefault((f, c), dict(fecha=f, cliente=c, planeados=0, por_enviar=0, sin_camion=0,
                                                 con_camion=0, despachados=0, solicitados=None))
            fila["solicitados"] = n
    out = []
    for (f, c), r in sorted(filas.items()):
        cubiertos = r["con_camion"] + r["despachados"]
        r["faltan"] = max(r["solicitados"] - cubiertos, 0) if r["solicitados"] is not None else None
        r["faltan_planear"] = max(r["solicitados"] - r["planeados"], 0) if r["solicitados"] is not None else None
        out.append(r)
    return out


# ---------------------------------------------------------------------------
# Destinos de prueba y mantenimiento
# ---------------------------------------------------------------------------
def cargar_destinos_prueba(get_conn_fn, cliente_dollar, cliente_diana, usuario):
    """Crea 5 tiendas 'Dollarcity 1..5' y 3 destinos 'Diana 1..3' (con coordenadas de prueba) en el catálogo de
    tiendas, y deja configurados los Cliente Orden GT-DOLLAR CITY y GT-DIANA si todavía no lo estaban.
    Devuelve (ok, mensaje)."""
    asignacion = {"GT-DOLLAR CITY": cliente_dollar, "GT-DIANA": cliente_diana}
    try:
        with closing(get_conn_fn()) as conn:
            try:
                with conn.cursor() as cur:
                    n = 0
                    for co, destinos in DESTINOS_PRUEBA.items():
                        cli = asignacion[co]
                        cur.execute("SELECT 1 FROM cat_clientes WHERE nombre = %s", (cli,))
                        if not cur.fetchone():
                            conn.rollback()
                            return False, f"El cliente '{cli}' no existe en Control de Ruta."
                        for nombre, codigo, lat, lon, _dir in destinos:
                            cur.execute("SELECT 1 FROM cat_clientes_tiendas WHERE cliente=%s AND tienda=%s", (cli, nombre))
                            if cur.fetchone():
                                continue
                            cur.execute("INSERT INTO cat_clientes_tiendas (cliente, tienda, km, codigo_tienda, clasificacion, "
                                        "latitud, longitud) VALUES (%s,%s,%s,%s,'Local',%s,%s)",
                                        (cli, nombre, 10.0, codigo, lat, lon))
                            n += 1
                        cur.execute("""INSERT INTO cat_transporte_cliente (cliente_orden, cliente, visit_type, activo)
                                       VALUES (%s,%s,%s,TRUE) ON CONFLICT (cliente_orden) DO NOTHING""",
                                    (co, cli, _prefijo_codigo(co).lower().replace(" ", "")))
                    _bitacora(cur, usuario, "destinos_prueba", detalle=f"{n} destino(s) de prueba creados")
                conn.commit()
            except Exception:
                conn.rollback()
                raise
        extra = []
        for co in DESTINOS_PRUEBA:
            ok_auto, dato = autoconfigurar_dummy(get_conn_fn, co, usuario)
            extra.append(f"{co}: dummy {dato}" if ok_auto else f"{co}: dummy sin identificar")
        return True, f"Listo: {n} destino(s) de prueba creados. " + "; ".join(extra) + "."
    except Exception as e:
        return False, f"No se pudieron crear los destinos de prueba ({type(e).__name__})."


def quitar_destinos_prueba(get_conn_fn, usuario):
    """Borra SOLO los destinos de prueba (por nombre y código exactos). Devuelve (ok, mensaje)."""
    try:
        with closing(get_conn_fn()) as conn:
            try:
                with conn.cursor() as cur:
                    n = 0
                    for destinos in DESTINOS_PRUEBA.values():
                        for nombre, codigo, *_ in destinos:
                            cur.execute("DELETE FROM cat_clientes_tiendas WHERE tienda = %s AND codigo_tienda = %s", (nombre, codigo))
                            n += cur.rowcount
                    _bitacora(cur, usuario, "quitar_destinos_prueba", detalle=f"{n} borrado(s)")
                conn.commit()
            except Exception:
                conn.rollback()
                raise
        return True, f"Se quitaron {n} destino(s) de prueba."
    except Exception as e:
        return False, f"No se pudieron quitar ({type(e).__name__}). Si ya se usaron en viajes, no se pueden borrar."


def purgar_planes_cerrados(get_conn_fn, usuario, dias=None):
    """Borra planes de más de `dias` días cuyos viajes ya están todos despachados o cancelados."""
    dias = dias or DIAS_RETENCION
    try:
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("""
                DELETE FROM plan_transporte p WHERE p.fecha < %s
                  AND NOT EXISTS (SELECT 1 FROM plan_viaje v WHERE v.plan_id = p.id
                                  AND v.estado NOT IN ('despachado','cancelado'))""",
                        (hoy_guatemala() - timedelta(days=dias),))
            n = cur.rowcount
            _bitacora(cur, usuario, "purgar", detalle=f"{n} plan(es) de más de {dias} días")
            conn.commit()
        return True, f"Se purgaron {n} plan(es) cerrados de más de {dias} días."
    except Exception as e:
        return False, f"No se pudo purgar ({type(e).__name__})."


# ===========================================================================
# PANTALLA (Streamlit). Se importa streamlit solo aquí adentro, para poder probar la lógica sin él.
# Reglas de la pantalla: nada pesado en session_state (solo ids, la vista previa de un archivo y mensajes),
# las consultas van acotadas por fecha/plan, y SR solo se consulta cuando se presiona un botón.
# ===========================================================================
import time as _time

_OPCIONES = {"t": 0.0, "placas": [], "piloto_de": {}, "pilotos": []}
_VEH_SR = {"t": 0.0, "lista": []}
_ICONO = {"Cancelado": "⚫", "Despachado": "🔵", "Con camión": "🟢", "En SR sin camión": "🟡"}


def _opciones_asignacion(get_conn_fn):
    """Camiones sincronizados con SR y pilotos activos. Una sola copia compartida, se refresca cada 2 minutos."""
    if _time.time() - _OPCIONES["t"] > 120 or not _OPCIONES["placas"]:
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("SELECT placa, piloto FROM cat_camiones WHERE id_sr IS NOT NULL AND COALESCE(activo, TRUE) ORDER BY placa")
            filas = cur.fetchall()
            cur.execute("SELECT DISTINCT dummy_prefijo FROM cat_transporte_cliente WHERE dummy_prefijo IS NOT NULL")
            prefs = [_clave_dummy(r[0]) for r in cur.fetchall()]
            filas = [f for f in filas if not any(re.match("^" + re.escape(pf) + r"\d+$", _clave_dummy(f[0])) for pf in prefs)
                     and not _clave_dummy(f[0]).startswith(DUMMY_GENERICO) and not es_generico(f[0])]
            cur.execute("SELECT nombre FROM cat_pilotos WHERE COALESCE(activo, TRUE) ORDER BY nombre")
            pil = [r[0] for r in cur.fetchall()]
        _OPCIONES.update(t=_time.time(), placas=[f[0] for f in filas], piloto_de={f[0]: f[1] for f in filas if f[1]}, pilotos=pil)
    return _OPCIONES


def diagnostico_vehiculos(get_conn_fn):
    """Resumen de los camiones del catálogo: cuántos se ofrecen para asignar y por qué no se ofrecen los demás."""
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT placa, activo, id_sr, COALESCE(pendiente_completar, FALSE) FROM cat_camiones")
        filas = cur.fetchall()
    r = {"disponibles": 0, "inactivos": [], "sin_sr": [], "genericos": [], "total": len(filas)}
    for placa, activo, id_sr, pend in filas:
        if id_sr is None:
            r["sin_sr"].append(placa)
        elif es_generico(placa):
            r["genericos"].append(placa)
        elif activo is False or pend:
            r["inactivos"].append(placa)
        else:
            r["disponibles"] += 1
    return r


def diagnosticar_placa(get_conn_fn, texto):
    """¿Por qué (no) aparece este camión? Devuelve (ok, mensaje)."""
    texto = _txt(texto)
    if not texto:
        return True, ""
    clave = sr_int.clave_placa(texto)
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT placa, activo, id_sr, COALESCE(pendiente_completar, FALSE) FROM cat_camiones")
        hallados = [r for r in cur.fetchall() if sr_int.clave_placa(r[0]) == clave or _clave_dummy(r[0]) == _clave_dummy(texto)]
    if not hallados:
        return False, (f"«{texto}» NO está en el catálogo de camiones de Control de Ruta. Aunque exista en SimpliRoute, hay que traerlo: "
                       "Catálogos → Integración SimpliRoute → sincronizar camiones.")
    placa, activo, id_sr, pend = hallados[0]
    if id_sr is None:
        return False, f"{placa} está en el catálogo pero NO tiene el id de SimpliRoute (no está sincronizado). Sincroniza camiones desde SimpliRoute."
    if es_generico(placa):
        return False, (f"{placa} se trata como vehículo GENÉRICO (su nombre es letras + número, no una placa), por eso nunca se ofrece como camión real. "
                       "Es correcto si es un dummy; si es un camión real con ese nombre, avísale a un Administrador.")
    if activo is False or pend:
        return False, (f"{placa} está sincronizado pero INACTIVO / pendiente de completar" + (" (lo creó la sincronización con SR)" if pend else "") +
                       ". Ábrelo en Catálogos → Camiones, complétalo y márcalo como activo; luego presiona «Actualizar lista» aquí.")
    return True, f"{placa} está disponible para asignar (id de SR {id_sr}). Si no lo ves, presiona «Actualizar lista»: la lista se guarda 2 minutos."


def _vehiculos_sr(token):
    """Lista ligera de vehículos de SR (id, nombre, placa, ubicación). Una sola copia, 5 minutos."""
    if _time.time() - _VEH_SR["t"] > 300 or not _VEH_SR["lista"]:
        ok, lista = sr_int.obtener_vehiculos_sr(token=token)
        if not ok:
            return False, lista
        _VEH_SR.update(t=_time.time(), lista=[dict(id=v.get("id"), nombre=v.get("name") or "", placa=v.get("license_plate") or "",
                                                  lat=v.get("location_start_latitude"), lon=v.get("location_start_longitude"),
                                                  dir=v.get("location_start_address") or "") for v in lista])
    return True, _VEH_SR["lista"]


def _flash(st, tipo, msg):
    st.session_state.setdefault("plan_flash", []).append((tipo, msg))


def _mostrar_flash(st):
    for tipo, msg in st.session_state.pop("plan_flash", []):
        getattr(st, tipo)(msg)


def _cada(st, clave, segundos):
    """True como máximo una vez cada `segundos` por sesión (para no repetir trabajo en cada clic)."""
    ahora = _time.time()
    if ahora - st.session_state.get(clave, 0) < segundos:
        return False
    st.session_state[clave] = ahora
    return True


def _reiniciar_widgets(st, viajes):
    for v in viajes:
        st.session_state.pop(f"plan_placa_{v['id']}", None)
        st.session_state.pop(f"plan_piloto_{v['id']}", None)


def render(st, get_conn_fn, usuario, perfil, clientes_permitidos, token_fn, error_fn=None, mostrar_error_fn=None):
    """Página 'Plan de Transporte'. clientes_permitidos: lista de clientes de Control de Ruta que el usuario puede ver."""
    import psycopg2

    st.header(":material/event_note: Plan de Transporte")
    if not PLAN_TRANSPORTE_ACTIVO:
        st.info("El módulo de Plan de Transporte está apagado por ahora.")
        return
    if perfil not in PERFILES_PLAN:
        st.warning("Tu perfil no tiene acceso a esta pantalla.")
        return
    if not asegurar_esquema_una_vez(get_conn_fn):
        st.error("No se pudo preparar la base de datos del Plan de Transporte. Avísale a un Administrador.")
        return
    if not tiene_acceso(get_conn_fn, perfil, clientes_permitidos):
        st.warning("No tienes acceso al Plan de Transporte: no tienes clientes de Transporte asignados. "
                   "Si lo necesitas, pídele a un Administrador que te los asigne.")
        return

    def _falla(e, ctx):
        if isinstance(e, (psycopg2.OperationalError, psycopg2.InterfaceError)):
            raise e                                   # lo maneja la pantalla general de "sin conexión"
        if error_fn and mostrar_error_fn:
            mostrar_error_fn(error_fn(e, ctx), perfil)
        else:
            st.error(f"❌ No se pudo completar ({type(e).__name__}).")

    cfgs = clientes_transporte(get_conn_fn, clientes_permitidos)
    _mostrar_flash(st)
    c1, c2 = st.columns([3, 2], vertical_alignment="center")
    c1.caption(f"Digitando como **{usuario}** ({perfil}). Aquí se cargan los planes de Transporte, se asignan los camiones y "
               "se ve el avance de todos tus clientes, sin entrar y salir de cada uno.")
    prueba = c2.checkbox("Modo prueba (el plan lleva ' -PRUEBA' en el nombre)", value=MODO_PRUEBA_POR_DEFECTO, key="plan_prueba",
                         help="Mientras esté marcado, todo lo que cargues va a planes de prueba, separados de los reales.")
    if prueba:
        st.warning("Modo prueba activo: los planes se crean como '… -PRUEBA' y no cuentan en el seguimiento real.", icon="🧪")

    ir_c = st.session_state.pop("plan_ir_carga", None)
    if ir_c:
        st.session_state["plan_sel_carga"] = ir_c
    ir_a = st.session_state.pop("plan_ir_a", None)      # tras guardar un plan: directo a asignar, en la fecha de ese plan
    if ir_a:
        st.session_state["plan_seccion"], st.session_state["plan_f_fecha_sel"] = ir_a
    secciones = ["Asignar camiones", "Cargar plan", "Seguimiento"] + (["Configuración"] if perfil in PERFILES_CONFIG else [])
    if not cfgs and perfil in PERFILES_CONFIG:
        st.session_state.setdefault("plan_seccion", "Configuración")     # primera vez: directo a configurar
        st.info("Primer paso: configura en esta pestaña los clientes de Transporte (cliente de Control de Ruta, vehículo dummy en SR y depósito).")
    sec = st.radio("Sección", secciones, horizontal=True, key="plan_seccion", label_visibility="collapsed")
    try:
        if sec == "Configuración":
            _seccion_config(st, get_conn_fn, usuario, token_fn, clientes_permitidos)
        elif not cfgs:
            st.info("Todavía no hay clientes de Transporte configurados para tu usuario. "
                    + ("Configúralos en la pestaña 'Configuración'." if perfil in PERFILES_CONFIG else "Pídele a un Administrador que los configure."))
        elif sec == "Asignar camiones":
            _seccion_asignar(st, get_conn_fn, usuario, token_fn, cfgs, prueba, _falla)
        elif sec == "Cargar plan":
            _seccion_cargar(st, get_conn_fn, usuario, cfgs, prueba, _falla, clientes_permitidos)
        else:
            _seccion_seguimiento(st, get_conn_fn, cfgs, clientes_permitidos, prueba)
    except Exception as e:
        _falla(e, f"plan_transporte_{sec}")


# ---------------------------------------------------------------------------
def _panel_plantilla(st, get_conn_fn, usuario, token_fn, prueba, clientes=None):
    """Indicador «X de Y planes enviados a SimpliRoute» de la plantilla (carga de Excel) elegida, con su botón de enviar."""
    import pandas as pd
    from datetime import timezone
    from zoneinfo import ZoneInfo
    cargas = listar_cargas(get_conn_fn, prueba, clientes=clientes)
    if not cargas:
        return
    ids = [c["id"] for c in cargas]
    if st.session_state.get("plan_sel_carga") not in ids:
        st.session_state.pop("plan_sel_carga", None)
    por_id = {c["id"]: c for c in cargas}

    def _etq(i):
        c = por_id[i]
        m = c["momento"].replace(tzinfo=timezone.utc).astimezone(ZoneInfo("America/Guatemala")) if c["momento"] else None
        return f"{c['archivo']}  ·  {m:%d/%m %H:%M}" if m else c["archivo"]
    sel = st.selectbox("Plantilla cargada", ids, key="plan_sel_carga", format_func=_etq)
    est = estado_carga(get_conn_fn, sel, clientes=clientes)
    if not est["planes"]:
        return
    total, env = est["planes_total"], est["planes_enviados"]
    pend_planes = [x for x in est["planes"] if not x["enviado"]]
    with st.container(border=True):
        (st.success if env == total else st.info)(f"**{env} de {total} planes enviados a SimpliRoute**  ·  "
                                                  f"{est['paradas_enviadas']} de {est['paradas_total']} paradas")
        st.progress(env / total if total else 0.0)
        st.dataframe(pd.DataFrame([{
            "Cliente": x["cliente"], "Fecha": x["fecha"].strftime("%a %d/%m"), "Plan": x["nombre"],
            "Estado": "✅ Enviado" if x["enviado"] else (f"⏳ {x['pendientes']} parada(s) por enviar" if x["pendientes"] else "⏳ Sin enviar"),
        } for x in est["planes"]]), hide_index=True, use_container_width=True)
        res = st.session_state.get("plan_resumen_carga")
        if res and res.get("carga") == sel:
            for f in res["filas"]:
                if f["errores"]:
                    st.warning(f"{f['cliente']} ({f['plan']}): {f['errores'][0]}")
                elif f["pendientes"]:
                    st.warning(f"{f['cliente']} ({f['plan']}): {f['mensaje']}")
        problemas = requisitos_envio(get_conn_fn, pend_planes, token_fn) if pend_planes else []
        for _av in (advertencias_envio(get_conn_fn, pend_planes) if pend_planes else []):
            st.warning(_av)
        if problemas:
            st.error("**Todavía no se puede enviar a SimpliRoute:**\n\n" + "\n".join(f"- {t}" for t in problemas))
        if st.button(f":material/cloud_upload: Enviar los {len(pend_planes)} plan(es) pendientes de esta plantilla" if pend_planes
                     else "Todos los planes de esta plantilla ya están enviados", key="plan_env_carga",
                     type="primary" if pend_planes else "secondary", disabled=not pend_planes or bool(problemas), use_container_width=True):
            token = token_fn()
            with st.spinner("Enviando los planes de la plantilla a SimpliRoute… (puede tardar un poco)"):
                resumen, _ok = enviar_planes(get_conn_fn, [x["id"] for x in pend_planes], usuario, token=token)
            st.session_state["plan_resumen_carga"] = {"carga": sel, "filas": resumen}
            st.rerun()


def _seccion_asignar(st, get_conn_fn, usuario, token_fn, cfgs, prueba, falla):
    hoy = hoy_guatemala()
    _panel_plantilla(st, get_conn_fn, usuario, token_fn, prueba, clientes=[c["cliente"] for c in cfgs])
    f1, f2 = st.columns([3, 1])
    opciones = ["Todos mis clientes"] + [c["cliente_orden"] for c in cfgs]
    sel = f1.selectbox("Cliente", opciones, key="plan_f_cliente")
    cos = [c["cliente_orden"] for c in cfgs] if sel == opciones[0] else [sel]
    fechas = fechas_con_plan(get_conn_fn, cos, prueba, hoy - timedelta(days=7))
    if not fechas:
        st.info(f"No hay planes{' de prueba' if prueba else ''} cargados desde hace una semana en adelante. "
                "Cárgalos en la sección 'Cargar plan'.")
        return
    if st.session_state.get("plan_f_fecha_sel") not in fechas:
        st.session_state.pop("plan_f_fecha_sel", None)
    proximas = [f for f in fechas if f >= hoy]
    kw = {} if "plan_f_fecha_sel" in st.session_state else {"index": fechas.index(proximas[0] if proximas else fechas[-1])}
    fecha = f2.selectbox("Fecha del plan", fechas, key="plan_f_fecha_sel", format_func=lambda d: d.strftime("%a %d/%m/%Y"), **kw)
    if _cada(st, "plan_sync_desp", 20):
        sincronizar_despachados(get_conn_fn)
    planes = listar_planes(get_conn_fn, clientes_orden=cos, fecha_ini=fecha, fecha_fin=fecha, prueba=prueba)
    if not planes:
        st.info(f"No hay planes{' de prueba' if prueba else ''} para el {fecha:%d/%m/%Y} con ese cliente.")
        return
    problemas_envio = _enviar_todos(st, get_conn_fn, usuario, token_fn, planes, fecha)
    if len(planes) > 1:
        idx = st.selectbox("Plan", range(len(planes)), format_func=lambda i: f"{planes[i]['nombre']}  ·  {planes[i]['viajes']} viaje(s)", key="plan_sel_plan")
        plan = planes[idx]
    else:
        plan = planes[0]
    d = detalle_plan(get_conn_fn, plan["id"])
    if not d:
        st.warning("El plan ya no existe.")
        return
    viajes = d["viajes"]
    if any(v.get("route_id_sr") and v["estado"] in ("en_sr", "asignado") for v in viajes) and _cada(st, f"plan_auto_ref_{plan['id']}", 120):
        _tok = token_fn()
        if _tok:       # una consulta ligera (rutas del día): completa camión/piloto que falten con lo que figure en SimpliRoute
            _ok_r, _msg_r, _camb = refrescar_desde_sr(get_conn_fn, plan["id"], usuario, token=_tok, solo_completar=True)
            if _camb:
                for _c in _camb[:10]:
                    _flash(st, "info", "Traído de SimpliRoute — " + _c)
                _reiniciar_widgets(st, viajes)
                st.rerun()
    activos = [v for v in viajes if v["estado"] != "cancelado"]
    con_camion = sum(1 for v in activos if v["estado"] in ("asignado", "despachado"))
    por_enviar_v, por_enviar_p = pendientes_de_envio(d)

    with st.container(border=True):
        st.markdown(f"**{plan['nombre']}**")
        if plan.get("plan_id_sr"):
            st.caption("Plan creado en SimpliRoute ✓")
        else:
            st.caption("Aún no se crea en SimpliRoute: se crea al enviar.")
        m1, m2, m3, m4 = st.columns(4)
        m1.metric("Viajes", len(activos))
        m2.metric("Con camión", con_camion)
        m3.metric("Faltan por asignar", len(activos) - con_camion)
        m4.metric("Por enviar a SR", por_enviar_v + por_enviar_p, help=f"{por_enviar_v} viaje(s) sin ruta en SR y {por_enviar_p} parada(s) sin enviar")
        if activos:
            st.progress(con_camion / len(activos), text=f"{con_camion} de {len(activos)} viajes con camión")
        b1, b2, b3 = st.columns([2, 2, 3])
        hay_pend = (por_enviar_v + por_enviar_p) > 0
        if b1.button(f":material/cloud_upload: Enviar pendientes a SR ({por_enviar_v + por_enviar_p})", key=f"plan_env_{plan['id']}",
                     type="primary" if hay_pend else "secondary", disabled=not hay_pend or bool(problemas_envio), use_container_width=True):
            token = token_fn()
            if not token:
                _flash(st, "error", "No hay token de SimpliRoute configurado.")
            else:
                with st.spinner("Enviando a SimpliRoute…"):
                    r = enviar_plan_a_sr(get_conn_fn, plan["id"], usuario, token=token)
                _flash(st, "success" if r["ok"] else "warning", r["mensaje"])
                for e in r["errores"][:8]:
                    _flash(st, "warning", e)
                for n in r["notas"][:5]:
                    _flash(st, "info", n)
                if r.get("excepcion") is not None:
                    falla(r["excepcion"], "plan_enviar_sr")
            st.rerun()
        if b2.button(":material/sync: Actualizar desde SR", key=f"plan_ref_{plan['id']}", use_container_width=True,
                     disabled=not any(v.get("route_id_sr") for v in activos)):
            if not _cada(st, f"plan_ref_t_{plan['id']}", 10):
                _flash(st, "info", "Se acaba de revisar contra SR; espera unos segundos.")
            else:
                token = token_fn()
                ok, msg, cambios = refrescar_desde_sr(get_conn_fn, plan["id"], usuario, token=token)
                _flash(st, "success" if ok else "warning", msg)
                for c in cambios[:10]:
                    _flash(st, "info", c)
                _reiniciar_widgets(st, viajes)
            st.rerun()
        with b3.popover("Vincular un plan que ya existe en SR", use_container_width=True):
            st.caption("SR no permite borrar planes. Si ya existe uno con este nombre, pega aquí su ID para reutilizarlo.")
            pid = st.text_input("ID del plan en SR", value=plan.get("plan_id_sr") or "", key=f"plan_vinc_{plan['id']}")
            if st.button("Vincular", key=f"plan_vinc_b_{plan['id']}"):
                ok, msg = vincular_plan_existente(get_conn_fn, plan["id"], pid, usuario)
                _flash(st, "success" if ok else "error", msg)
                st.rerun()

    if not viajes:
        st.info("Este plan todavía no tiene viajes.")
        return
    with st.expander("¿No aparece un camión para asignar?"):
        dg = diagnostico_vehiculos(get_conn_fn)
        g1, g2, g3, g4 = st.columns(4)
        g1.metric("Disponibles", dg["disponibles"], help="Activos, con id de SimpliRoute y que no son genéricos.")
        g2.metric("Inactivos / por completar", len(dg["inactivos"]), help="Sincronizados desde SR pero sin completar o inactivos en Catálogos.")
        g3.metric("Sin id de SR", len(dg["sin_sr"]), help="Están en el catálogo pero no sincronizados con SimpliRoute.")
        g4.metric("Genéricos (dummies)", len(dg["genericos"]), help="Nombre de letras + número (DC01, ADICIONALES1…): no se ofrecen como camión real.")
        if dg["inactivos"]:
            st.caption("Inactivos / por completar: " + ", ".join(dg["inactivos"][:30]) + (" …" if len(dg["inactivos"]) > 30 else ""))
        busca = st.text_input("Escribe la placa que no encuentras", key="plan_diag_placa", placeholder="ej. 330BHZ o C-330BHZ")
        if busca.strip():
            ok_d, msg_d = diagnosticar_placa(get_conn_fn, busca)
            (st.success if ok_d else st.warning)(msg_d)
        if st.button("Actualizar lista de camiones", key="plan_diag_refrescar"):
            _OPCIONES["t"] = 0.0
            st.rerun()
    opts = _opciones_asignacion(get_conn_fn)
    for v in viajes:
        _tarjeta_viaje(st, get_conn_fn, usuario, token_fn, plan, v, opts, falla)


def _enviar_todos(st, get_conn_fn, usuario, token_fn, planes, fecha):
    """Botón principal: sube a SR TODOS los planes de esa fecha (de los clientes filtrados) y dice cuántos subió de cada cliente.
    Devuelve la lista de requisitos que faltan (vacía si se puede enviar)."""
    import pandas as pd
    problemas = requisitos_envio(get_conn_fn, planes, token_fn)
    if problemas:
        st.error("**Todavía no se puede enviar a SimpliRoute:**\n\n" + "\n".join(f"- {t}" for t in problemas))
    res = st.session_state.get("plan_resumen_envio")
    if res and res.get("fecha") == fecha:
        filas = res["filas"]
        ok_n = sum(1 for f in filas if not f["pendientes"] and not f["errores"] and (f["rutas"] or f["paradas"] or "Ya estaba" in f["mensaje"]))
        caja = st.success if ok_n == len(filas) else st.warning
        caja(f"Resultado del último envío: {ok_n} de {len(filas)} plan(es) completos en SimpliRoute."
             + ("" if ok_n == len(filas) else " Revisa los que tienen pendientes o errores."))
        st.dataframe(pd.DataFrame([{
            "Cliente": f["cliente"], "Plan": f["plan"],
            "Plan en SR": "Creado ahora" if f["plan_nuevo_en_sr"] else "Ya existía",
            "Rutas enviadas": str(f["rutas"]), "Visitas enviadas": str(f["paradas"]), "Pendientes": str(f["pendientes"]),
            "Detalle": (f["errores"][0] if f["errores"] else f["mensaje"])} for f in filas]), hide_index=True, use_container_width=True)
        for f in filas:
            for e in f["errores"][1:6]:
                st.caption(f"⚠️ {f['cliente']}: {e}")
        if st.button("Ocultar resultado", key="plan_ocultar_res"):
            st.session_state.pop("plan_resumen_envio", None)
            st.rerun()
    pend_planes = 0
    for pl in planes:
        d = detalle_plan(get_conn_fn, pl["id"])
        vp, pp = pendientes_de_envio(d) if d else (0, 0)
        if vp + pp:
            pend_planes += 1
    etiqueta = (f":material/cloud_upload: Enviar todos los planes del {fecha:%d/%m} a SR ({pend_planes} plan(es) con pendientes)"
                if pend_planes else f"Todos los planes del {fecha:%d/%m} ya están enviados a SR")
    if st.button(etiqueta, key="plan_env_todos", type="primary" if pend_planes else "secondary",
                 disabled=not pend_planes or bool(problemas), use_container_width=True,
                 help="Sube a SimpliRoute todos los planes de esta fecha, uno por cliente, y al final te dice cuántos subió de cada uno."):
        token = token_fn()
        with st.spinner("Enviando los planes a SimpliRoute… (puede tardar un poco)"):
            resumen, _ok = enviar_planes(get_conn_fn, [pl["id"] for pl in planes], usuario, token=token)
        st.session_state["plan_resumen_envio"] = {"fecha": fecha, "filas": resumen}
        st.rerun()
    return problemas


def _tarjeta_viaje(st, get_conn_fn, usuario, token_fn, plan, v, opts, falla):
    est = estado_visible(v)
    cerrado = v["estado"] in ("despachado", "cancelado")
    paradas_act = [p for p in v["paradas"] if p["estado"] != "cancelada"]
    with st.container(border=True):
        c0, c1, c2, c3, c4 = st.columns([3.4, 1.9, 1.9, 1.2, 0.7], vertical_alignment="top")
        titulo = f"**Viaje {v['num']}**" + (f" · grupo {v['grupo']}" if v["grupo"] is not None else "")
        if len(paradas_act) == 1:
            titulo += f" · {paradas_act[0]['destino']}"          # un solo destino: va en el mismo renglón
        c0.markdown(f"{titulo}  \n{_ICONO.get(est.split(' ·')[0], '⚪')} {est}")
        if len(paradas_act) > 1:                                  # varios destinos: todos a la vista, sin abrir nada
            for p in paradas_act:
                c0.caption(f"{p['orden']}. {p['destino']}  ·  {p['codigo_pedido']}" + ("  ⚠️" if p.get("error") else ""))
        elif not paradas_act:
            c0.caption("(sin paradas activas)")
        if cerrado:
            c1.write(v.get("placa") or "—")
            c2.write(v.get("piloto") or "—")
            if v["estado"] == "despachado" and v.get("id_viaje_cr"):
                c3.caption(f"Viaje {v['id_viaje_cr']}")
        else:
            lp = [""] + opts["placas"]
            lpi = [""] + opts["pilotos"]
            kp, kpi = f"plan_placa_{v['id']}", f"plan_piloto_{v['id']}"
            if kp not in st.session_state:
                st.session_state[kp] = v["placa"] if v.get("placa") in lp else ""
            if kpi not in st.session_state:
                st.session_state[kpi] = v["piloto"] if v.get("piloto") in lpi else ""

            def _cb(vid=v["id"]):
                pl = st.session_state.get(f"plan_placa_{vid}")
                sug = _OPCIONES["piloto_de"].get(pl)
                if pl and sug and sug in _OPCIONES["pilotos"]:
                    st.session_state[f"plan_piloto_{vid}"] = sug
            c1.selectbox("Placa", lp, key=kp, label_visibility="collapsed", on_change=_cb,
                         format_func=lambda x: x or "— Placa —")
            c2.selectbox("Piloto", lpi, key=kpi, label_visibility="collapsed", format_func=lambda x: x or "— Piloto —")
            cambio = (st.session_state[kp] or None) != v.get("placa") or (st.session_state[kpi] or None) != v.get("piloto")
            if c3.button("Asignar", key=f"plan_asig_{v['id']}", type="primary" if cambio else "secondary",
                         disabled=not st.session_state[kp], use_container_width=True):
                token = token_fn() if v.get("route_id_sr") else None
                ok, msg, avisos = asignar_viaje(get_conn_fn, v["id"], st.session_state[kp], st.session_state[kpi] or None, usuario, token=token)
                _flash(st, "success" if ok else "error", f"Viaje {v['num']}: {msg}")
                for a in avisos:
                    _flash(st, "warning", a)
                st.rerun()
            with c4.popover("⋯", help="Opciones del viaje", use_container_width=True):
                for p in v["paradas"]:
                    if p["estado"] == "cancelada":
                        continue
                    x, y = st.columns([3, 1.4], vertical_alignment="center")
                    x.write(f"{p['orden']}. {p['destino']}  \n:gray[{p['codigo_pedido']} · "
                            + {"pendiente": "por enviar", "enviando": "por verificar", "en_sr": "en SR", "error": "con error"}.get(p["estado"], p["estado"]) + "]")
                    if y.button("Quitar", key=f"plan_qp_{p['id']}", use_container_width=True):
                        ok, msg = cancelar_parada(get_conn_fn, p["id"], usuario, token=token_fn() if p.get("visit_id_sr") else None)
                        _flash(st, "success" if ok else "error", msg)
                        _reiniciar_widgets(st, [v])
                        st.rerun()
                if st.button("Cancelar todo el viaje", key=f"plan_cvb_{v['id']}", use_container_width=True):
                    st.session_state[f"plan_cv_{v['id']}"] = True
                    st.rerun()
        if v.get("cambiado_en_sr"):
            st.caption("🔁 Este viaje se cambió directamente en SimpliRoute; Control de Ruta ya lo tomó.")
        if v.get("error"):
            st.caption(f"⚠️ {_corto(v['error'], 200)}")
        for p in paradas_act:
            if p.get("error") and len(paradas_act) > 1:
                st.caption(f"⚠️ {p['codigo_pedido']}: {_corto(p['error'], 160)}")
        kc = f"plan_cv_{v['id']}"
        if not cerrado and st.session_state.get(kc):
            st.warning("¿Cancelar el viaje completo? Si ya está en SimpliRoute, se borran su ruta y sus visitas allá.")
            s1, s2 = st.columns(2)
            if s1.button("Sí, cancelar el viaje", key=f"plan_cvs_{v['id']}", type="primary"):
                ok, msg = cancelar_viaje(get_conn_fn, v["id"], usuario, token=token_fn() if v.get("route_id_sr") else None)
                st.session_state.pop(kc, None)
                _flash(st, "success" if ok else "error", f"Viaje {v['num']}: {msg}")
                st.rerun()
            if s2.button("No", key=f"plan_cvn_{v['id']}"):
                st.session_state.pop(kc, None)
                st.rerun()


# ---------------------------------------------------------------------------
def _seccion_cargar(st, get_conn_fn, usuario, cfgs, prueba, falla, clientes_permitidos=None):
    import hashlib
    import pandas as pd

    st.markdown("**Sube la plantilla de carga masiva** (hoja PEDIDOS). Cada fila es una parada; las filas con el mismo "
                "Cliente Orden, fecha y **G. Ruta** forman un solo viaje. Si G. Ruta va vacía, la fila es su propio viaje.")
    n_up = st.session_state.get("plan_up_n", 0)
    archivo = st.file_uploader("Plantilla Excel (.xlsx)", type=["xlsx"], key=f"plan_up_{n_up}")
    permitir = st.checkbox("Permitir paradas repetidas (mismo destino en el mismo viaje y fecha)", key="plan_permitir_rep")
    if archivo is not None:
        datos = archivo.getvalue()
        sig = (hashlib.md5(datos).hexdigest(), bool(prueba), bool(permitir))
        if st.session_state.get("plan_up_sig") != sig:
            filas, errs_glob = leer_excel_plan(datos)
            val = validar_filas(get_conn_fn, filas, prueba=prueba, permitir_repetidas=permitir, clientes_permitidos=clientes_permitidos) if filas else []
            st.session_state.update(plan_up_sig=sig, plan_up_val=val, plan_up_errs=errs_glob)
        val, errs_glob = st.session_state.get("plan_up_val", []), st.session_state.get("plan_up_errs", [])
        for e in errs_glob:
            st.error(e)
        if val:
            ok_f = [f for f in val if f["ok"]]
            mal = [f for f in val if not f["ok"]]
            m1, m2, m3 = st.columns(3)
            m1.metric("Filas leídas", len(val))
            m2.metric("Válidas", len(ok_f))
            m3.metric("Con error", len(mal))
            if mal and not ok_f and all(any("Ya está cargada" in e for e in f["errores"]) for f in val) and any(f["fecha"] for f in val):
                st.info("Este archivo ya está guardado como plan. Ve a «Asignar camiones» para enviarlo a SimpliRoute y asignar los camiones.")
                if st.button("Ir a Asignar camiones", key="plan_ir_asig"):
                    st.session_state["plan_ir_a"] = ("Asignar camiones", min(f["fecha"] for f in val if f["fecha"]))
                    st.rerun()
            if ok_f and not prueba and any(f["fecha"] <= hoy_guatemala() for f in ok_f):
                st.warning("Estás cargando un plan para **hoy** en modo REAL: al enviarlo se crearán rutas reales en SimpliRoute que se mezclan con las "
                           "de hoy (y que «Importar Rutas» de Despacho verá). Para un piloto, marca «Modo prueba» o usa una fecha futura.")
            if ok_f:
                st.markdown("**Se crearía:**")
                for (co, fe), r in agrupar_para_vista(ok_f).items():
                    st.write(f"• `{nombre_plan(co, fe, prueba)}` — {r['viajes']} viaje(s) nuevo(s) y {r['paradas']} parada(s)")
            solo_validas = True
            if mal:
                solo_validas = st.checkbox("Guardar solo las filas válidas (las que tienen error se ignoran)", key="plan_solo_validas")
            if st.button(":material/save: Guardar plan", type="primary", key="plan_guardar", disabled=not ok_f or not solo_validas,
                         help="Queda guardado como borrador en Control de Ruta; se envía a SimpliRoute desde 'Asignar camiones'."):
                r = guardar_filas(get_conn_fn, ok_f, usuario, prueba=prueba, archivo=getattr(archivo, "name", None))
                if r["ok"]:
                    st.session_state["plan_ir_carga"] = r.get("carga_id")
                    for pl in r["planes"]:
                        extra = f", {pl['a_viaje_existente']} agregada(s) a un viaje existente" if pl["a_viaje_existente"] else ""
                        _flash(st, "success", f"{pl['nombre']}: {pl['paradas']} parada(s), {pl['viajes_nuevos']} viaje(s) nuevo(s){extra}.")
                    _flash(st, "info", "Plan guardado. Aquí lo ves para enviarlo a SimpliRoute y asignar camiones.")
                    for k in ("plan_up_sig", "plan_up_val", "plan_up_errs", "plan_solo_validas"):
                        st.session_state.pop(k, None)
                    st.session_state["plan_up_n"] = n_up + 1          # vacía el cargador y libera el archivo
                    st.session_state["plan_ir_a"] = ("Asignar camiones", min(f["fecha"] for f in ok_f))
                    st.rerun()
                else:
                    st.error(r["mensaje"])
                    if r.get("excepcion") is not None:
                        falla(r["excepcion"], "plan_guardar")
            df = pd.DataFrame([{
                "Fila": str(f["n"]), "Cliente": str(f["cliente_orden"] or ""), "Fecha": f["fecha"].strftime("%d/%m/%Y") if f["fecha"] else "",
                "G. Ruta": str(f["grupo"]) if f["grupo"] is not None else "", "Destino": str(f["destino"] or ""),
                "Resultado": "✅ Válida" if f["ok"] else "❌ Con error",
                "Detalle": " · ".join(f["errores"] + f["avisos"])} for f in val])
            st.dataframe(df, hide_index=True, use_container_width=True)
    else:
        for k in ("plan_up_sig", "plan_up_val", "plan_up_errs"):
            st.session_state.pop(k, None)

    with st.expander(":material/add_circle: Agregar una parada a mano (a un plan nuevo o ya existente)"):
        co_opts = [c["cliente_orden"] for c in cfgs]
        a1, a2 = st.columns(2)
        co = a1.selectbox("Cliente Orden", co_opts, key="plan_m_co")
        fe = a2.date_input("Fecha de entrega", value=hoy_guatemala(), min_value=hoy_guatemala(), key="plan_m_fe", format="DD/MM/YYYY")
        cli = next(c["cliente"] for c in cfgs if c["cliente_orden"] == co)
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("SELECT tienda, codigo_tienda FROM cat_clientes_tiendas WHERE cliente = %s ORDER BY tienda LIMIT 500", (cli,))
            tiendas = cur.fetchall()
        if not tiendas:
            st.info("Este cliente no tiene destinos en el catálogo. Sube el Excel con las coordenadas, o carga destinos de prueba en 'Configuración'.")
        else:
            b1, b2 = st.columns([3, 1])
            ti = b1.selectbox("Destino", range(len(tiendas)), key="plan_m_dest",
                              format_func=lambda i: f"{tiendas[i][1]} · {tiendas[i][0]}" if tiendas[i][1] is not None else tiendas[i][0])
            gr = b2.number_input("G. Ruta (0 = viaje propio)", min_value=0, step=1, value=0, key="plan_m_gr")
            gens = [v["nombre"] for v in vehiculos_sincronizados(get_conn_fn) if es_generico(v["nombre"])]
            idx_g = next((i for i, n in enumerate(gens) if _clave_dummy(n).startswith(DUMMY_GENERICO)), 0)
            veh_m = st.selectbox("Vehículo (Placa)", gens, index=idx_g if gens else None, key="plan_m_veh",
                                 help="El genérico con el que se crea la ruta en SimpliRoute (como DC01 o ADICIONALES1). El camión real se asigna después.") if gens else None
            if not gens:
                st.caption("No encontré vehículos genéricos sincronizados; sincroniza camiones desde SimpliRoute.")
            obs = st.text_input("Observaciones (opcional)", key="plan_m_obs")
            if st.button("Agregar parada", key="plan_m_btn"):
                tienda, codigo = tiendas[ti]
                fila = fila_manual(co, fe, str(codigo) if codigo is not None else tienda, tienda, grupo=int(gr) or None, obs=obs, placa=veh_m or "")
                v = validar_filas(get_conn_fn, [fila], prueba=prueba, permitir_repetidas=permitir, clientes_permitidos=clientes_permitidos)
                if not v[0]["ok"]:
                    for e in v[0]["errores"]:
                        st.error(e)
                else:
                    r = guardar_filas(get_conn_fn, v, usuario, prueba=prueba)
                    if r["ok"]:
                        pl = r["planes"][0]
                        _flash(st, "success", f"Parada agregada a {pl['nombre']}" + (" (a un viaje existente)." if pl["a_viaje_existente"] else "."))
                        st.rerun()
                    else:
                        st.error(r["mensaje"])


# ---------------------------------------------------------------------------
def _seccion_seguimiento(st, get_conn_fn, cfgs, clientes_permitidos, prueba):
    import pandas as pd
    hoy = hoy_guatemala()
    a, b = st.columns(2)
    ini = a.date_input("Desde", value=hoy, key="plan_s_ini", format="DD/MM/YYYY")
    fin = b.date_input("Hasta", value=hoy + timedelta(days=6), key="plan_s_fin", format="DD/MM/YYYY")
    if fin < ini or (fin - ini).days > 62:
        st.warning("Elige un rango válido de hasta 62 días.")
        return
    nombres = sorted({c["cliente"] for c in cfgs})
    filas = resumen_seguimiento(get_conn_fn, ini, fin, clientes=nombres, prueba=prueba)
    if not filas:
        st.info("No hay planes ni solicitudes de vehículos en ese rango.")
        return
    tot = lambda k: sum(f[k] or 0 for f in filas)
    m = st.columns(5)
    m[0].metric("Solicitados", tot("solicitados"))
    m[1].metric("Planeados", tot("planeados"))
    m[2].metric("Con camión", tot("con_camion") + tot("despachados"))
    m[3].metric("Faltan con camión", tot("faltan"))
    m[4].metric("Faltan por planear", tot("faltan_planear"))
    df = pd.DataFrame([{
        "Fecha": f["fecha"].strftime("%a %d/%m"), "Cliente": f["cliente"],
        "Solicitados": "—" if f["solicitados"] is None else str(f["solicitados"]), "Planeados": str(f["planeados"]),
        "Por enviar a SR": str(f["por_enviar"]), "En SR sin camión": str(f["sin_camion"]), "Con camión": str(f["con_camion"]),
        "Despachados": str(f["despachados"]), "Faltan": "—" if f["faltan"] is None else str(f["faltan"]),
        "Faltan por planear": "—" if f["faltan_planear"] is None else str(f["faltan_planear"])} for f in filas])
    st.dataframe(df, hide_index=True, use_container_width=True)
    st.caption("Faltan = solicitados menos los viajes que ya tienen camión o ya se despacharon. «—» = ese día no hay solicitud registrada.")


# ---------------------------------------------------------------------------
def _seccion_config(st, get_conn_fn, usuario, token_fn, clientes_permitidos):
    import pandas as pd
    st.markdown("**Clientes de Transporte** — qué cliente de Control de Ruta corresponde a cada «Cliente Orden» de la plantilla, "
                "su vehículo dummy en SR y su depósito de salida.")
    cfgs = listar_config(get_conn_fn)
    if cfgs:
        st.dataframe(pd.DataFrame([{"Cliente Orden": c["cliente_orden"], "Cliente": c["cliente"], "visit_type": c["visit_type"] or "",
                                    "Dummies": (f"grupo {c['dummy_prefijo']}" if c.get("dummy_prefijo") else (c["vehiculo_dummy_nombre"] or f"genéricos {DUMMY_GENERICO}")), "Depósito": c["deposito_nombre"] or "",
                                    "Activo": "Sí" if c["activo"] else "No"} for c in cfgs]), hide_index=True, use_container_width=True)
    with closing(get_conn_fn()) as conn, conn.cursor() as cur:
        cur.execute("SELECT nombre FROM cat_clientes WHERE COALESCE(activo, TRUE) ORDER BY nombre")
        clientes_cr = [r[0] for r in cur.fetchall()]
    with st.expander(":material/edit: Agregar o editar un cliente de Transporte", expanded=not cfgs):
        co = st.selectbox("Cliente Orden (de la plantilla)", CLIENTES_ORDEN_PARAMETROS, key="plan_c_co")
        actual = next((c for c in cfgs if c["cliente_orden"] == co), {})
        a1, a2 = st.columns(2)
        cli = a1.selectbox("Cliente en Control de Ruta", clientes_cr, key="plan_c_cli",
                           index=clientes_cr.index(actual["cliente"]) if actual.get("cliente") in clientes_cr else 0) if clientes_cr else None
        vt = a2.text_input("visit_type en SR", value=actual.get("visit_type") or _prefijo_codigo(co).lower().replace(" ", ""), key="plan_c_vt",
                           help="Así reconoce Despacho a qué cliente pertenecen las rutas importadas de SR.")
        st.caption("El vehículo de cada viaje viene en la columna **Placa** del Excel; aquí no hay que elegir dummies.")
        # variables por defecto (el bloque opcional las redefine solo si se abre y se usa)
        grupo_sel, un_solo, vsel, vehs, en_sr = None, False, None, [], False
        with st.expander("Opcional: vehículo por defecto para filas sin Placa"):
            st.markdown("**Dummies del cliente**")
            grupos = grupos_dummy(get_conn_fn)
            gen = dummies_genericos(get_conn_fn)
            st.caption(f"Si no eliges nada, el cliente usa los genéricos «{DUMMY_GENERICO}» que ya están sincronizados "
                       + (f"({len(gen)}: {gen[0]['nombre']} … {gen[-1]['nombre']}). " if gen else "(no encontré ninguno todavía). ")
                       + "Solo elige un grupo si el cliente tiene el suyo (como DC01, DC02…). No se crea ningún vehículo: cada viaje toma el siguiente dummy.")
            sug = sugerir_prefijo(co, grupos)
            prefs = [""] + sorted(grupos)
            ya = _clave_dummy(actual.get("dummy_prefijo") or "")
            idx_def = prefs.index(ya) if ya in prefs else 0        # por defecto ninguno: sin Placa en el Excel = error de plantilla
            grupo_sel = st.selectbox("Grupo de dummies", prefs, index=idx_def, key=f"plan_c_grp_{co}",
                                     format_func=lambda x: f"Genéricos «{DUMMY_GENERICO}» (por defecto)" if not x else
                                     f"{x}  ·  {len(grupos[x])} vehículo(s): {grupos[x][0]['nombre']} … {grupos[x][-1]['nombre']}")
            if sug and not ya:
                st.caption(f"Sugerido por las iniciales de «{co}»: **{sug}**.")
            if not grupos:
                st.caption("No encontré dummies sincronizados. Sincroniza camiones desde SimpliRoute en Catálogos → Integración SimpliRoute.")
            un_solo = st.checkbox("Este cliente usa un solo vehículo dummy (no un grupo)", key=f"plan_c_uno_{co}", value=bool(actual.get("vehiculo_dummy_sr") and not ya))
            vsel, vehs, en_sr = None, [], False
            if un_solo:
                if actual.get("vehiculo_dummy_nombre"):
                    st.caption(f"Actual: {actual['vehiculo_dummy_nombre']} (id {actual['vehiculo_dummy_sr']})")
                en_sr = st.checkbox("No aparece: buscarlo directo en SimpliRoute", key="plan_c_fsr")
                filtro = st.text_input("Buscar por nombre", key="plan_c_vf", placeholder="ej. unisuper")
                if en_sr:
                    if st.button("Cargar vehículos de SR", key="plan_c_vb"):
                        ok, lista = _vehiculos_sr(token_fn())
                        if not ok:
                            st.error(f"No pude consultar SR: {lista}")
                    vehs = [dict(id=v["id"], nombre=v["nombre"] or v["placa"], lat=v["lat"], lon=v["lon"]) for v in _VEH_SR["lista"]]
                else:
                    vehs = vehiculos_sincronizados(get_conn_fn)
                vehs = [v for v in vehs if not filtro or filtro.lower() in str(v["nombre"]).lower()][:200]
                vsel = st.selectbox("Vehículo dummy", range(len(vehs)), key="plan_c_vs",
                                    format_func=lambda i: f"{vehs[i]['nombre'] or '(sin nombre)'} · id {vehs[i]['id']}") if vehs else None
        d1, d2, d3 = st.columns([3, 1, 1])
        dep = d1.selectbox("Depósito de salida y retorno", [""] + DEPOSITOS_PARAMETROS, key="plan_c_dep",
                           index=([""] + DEPOSITOS_PARAMETROS).index(actual["deposito_nombre"]) if actual.get("deposito_nombre") in DEPOSITOS_PARAMETROS else 0)
        base = sr_int.UBICACION_BASE_VEHICULOS_SR
        lat = d2.number_input("Latitud", value=float(actual.get("deposito_lat") or base["latitud"]), format="%.6f", key="plan_c_lat")
        lon = d3.number_input("Longitud", value=float(actual.get("deposito_lon") or base["longitud"]), format="%.6f", key="plan_c_lon")
        usar_dummy = st.checkbox("Tomar las coordenadas del depósito de la ubicación del vehículo (solo si lo buscaste en SR)", key="plan_c_ud") if en_sr else False
        act = st.checkbox("Activo", value=actual.get("activo", True), key="plan_c_act")
        if st.button("Guardar configuración", type="primary", key="plan_c_save", disabled=not cli):
            datos = dict(cliente_orden=co, cliente=cli, visit_type=vt, deposito_nombre=dep or None, deposito_lat=lat, deposito_lon=lon, activo=act,
                         vehiculo_dummy_sr=None, vehiculo_dummy_nombre=None, dummy_prefijo=None)
            if un_solo:
                if vsel is not None:
                    v = vehs[vsel]
                    datos.update(vehiculo_dummy_sr=v["id"], vehiculo_dummy_nombre=v["nombre"])
                    if usar_dummy and v.get("lat") and v.get("lon"):
                        datos.update(deposito_lat=float(v["lat"]), deposito_lon=float(v["lon"]))
                else:
                    datos.update(vehiculo_dummy_sr=actual.get("vehiculo_dummy_sr"), vehiculo_dummy_nombre=actual.get("vehiculo_dummy_nombre"))
            else:
                datos.update(dummy_prefijo=grupo_sel or None)
            ok, msg = guardar_config(get_conn_fn, datos, usuario)
            (st.success if ok else st.error)(msg)
            if ok:
                st.rerun()

    with st.expander(":material/science: Destinos de prueba (Dollarcity 1 a 5 y Diana 1 a 3)"):
        st.caption("Crea 5 tiendas «Dollarcity 1..5» y 3 destinos «Diana 1..3» (códigos 9001–9005 y 9101–9103, con coordenadas de prueba) "
                   "y deja configurados los Cliente Orden GT-DOLLAR CITY y GT-DIANA si faltaban. Después puedes quitarlos.")
        def _pref(txt):
            for i, c in enumerate(clientes_cr):
                if txt in c.lower():
                    return i
            return 0
        p1, p2 = st.columns(2)
        cd = p1.selectbox("Cliente de Control de Ruta para Dollarcity", clientes_cr, index=_pref("dollar"), key="plan_t_cd") if clientes_cr else None
        cx = p2.selectbox("Cliente de Control de Ruta para Diana", clientes_cr, index=_pref("diana"), key="plan_t_cx") if clientes_cr else None
        q1, q2 = st.columns(2)
        if q1.button("Crear destinos de prueba", key="plan_t_crear", disabled=not (cd and cx)):
            ok, msg = cargar_destinos_prueba(get_conn_fn, cd, cx, usuario)
            (st.success if ok else st.error)(msg)
        confirma = q2.checkbox("Confirmo que quiero quitarlos", key="plan_t_conf")
        if q2.button("Quitar destinos de prueba", key="plan_t_quitar", disabled=not confirma):
            ok, msg = quitar_destinos_prueba(get_conn_fn, usuario)
            (st.success if ok else st.error)(msg)

    with st.expander(":material/cleaning_services: Mantenimiento y bitácora"):
        rep_cod = codigos_repetidos(get_conn_fn)
        if rep_cod:
            st.warning(f"Hay {len(rep_cod)} código(s) repetido(s) dentro de un mismo cliente. Un código debe identificar UNA sola tienda; "
                       "corrígelos en Catálogos → Clientes y Tiendas (mientras existan, la base no puede bloquear nuevos repetidos): "
                       + "; ".join(f"{c} · código {k}: {', '.join(t)}" for c, k, t in rep_cod[:8]) + (" …" if len(rep_cod) > 8 else ""))
        cdias, cbtn = st.columns([1, 3], vertical_alignment="bottom")
        dias_sr = cdias.number_input("Días de SR a revisar", min_value=1, max_value=14, value=DIAS_COORD_SR, key="plan_dias_coord")
        if cbtn.button("Buscar en SimpliRoute las coordenadas que falten en el catálogo", key="plan_coord_sr",
                       help="Copia a Catálogos la latitud y longitud de las tiendas que aparecen en las visitas recientes de SR (cliente + código). "
                            "No toca nombre ni km. Cada día revisado tarda unos 15 segundos."):
            with st.spinner("Consultando SimpliRoute…"):
                ok, resumen, detalle = sr_int.sincronizar_coordenadas_tiendas_desde_sr(get_conn_fn, dias=int(dias_sr), token=token_fn())
            (st.success if ok else st.error)(resumen)
            if detalle:
                st.caption(detalle)
        if st.button(f"Purgar planes cerrados de más de {DIAS_RETENCION} días", key="plan_purgar"):
            ok, msg = purgar_planes_cerrados(get_conn_fn, usuario)
            (st.success if ok else st.error)(msg)
        with closing(get_conn_fn()) as conn, conn.cursor() as cur:
            cur.execute("SELECT momento, usuario, accion, detalle FROM plan_bitacora ORDER BY id DESC LIMIT 40")
            filas = cur.fetchall()
        if filas:
            st.dataframe(pd.DataFrame(filas, columns=["Momento", "Usuario", "Acción", "Detalle"]), hide_index=True, use_container_width=True)
